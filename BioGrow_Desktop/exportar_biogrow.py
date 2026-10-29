"""
exportar_biogrow.py — Guarda varias gráficas en un solo archivo (PDF o Excel).

- PDF: una portada con la información del proyecto y el resultado del ANOVA, y luego una gráfica por página.
- Excel: hoja 'Proyecto', hoja 'Gráficas' con las imágenes y una hoja por tabla de cálculos
  (ranking, promedio/DE/varianza, ANOVA, Tukey, LSD) más los datos de cada grupo.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import List, Optional, Tuple

import pandas as pd

import analisis_biogrow as AN

VERDE = "059669"


def _png(fig, dpi=150) -> Tuple[bytes, float, float]:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, facecolor=fig.get_facecolor(), bbox_inches="tight")
    from PIL import Image as PILImage
    w, h = PILImage.open(io.BytesIO(buf.getvalue())).size  # tamaño real tras recortar márgenes
    return buf.getvalue(), w, h


def _dictamen(res: dict) -> str:
    an = res.get("anova")
    if not an:
        return "Sin ANOVA (se necesitan al menos 2 grupos con 2 o más datos)."
    return (("✔ Sí hay" if an["significativo"] else "✘ No hay") +
            f" diferencias significativas en {res['nombre_var'].lower()} entre los grupos "
            f"(F calculada = {an['F calculada']:.3f} {'>' if an['significativo'] else '≤'} F crítica = "
            f"{an['F crítica']:.3f}; GL = {an['GL entre']} y {an['GL error']}; p = {an['p-valor']:.4f}; α = {res['alfa']}).")


def tablas_calculos(res: dict) -> List[Tuple[str, pd.DataFrame]]:
    if not res or not res.get("ok"):
        return []
    u = res["u"]
    salida = [("Ranking", res["ranking"]), ("Promedio DE varianza", res["tabla_descriptivos"])]
    if res.get("anova"):
        salida.append(("ANOVA", res["tabla_anova"]))
    if res.get("tabla_tukey") is not None and not res["tabla_tukey"].empty:
        salida.append(("Tukey", res["tabla_tukey"]))
        salida.append(("Letras Tukey", res["letras_tukey"].rename(columns={"Media": f"Promedio ({u})"})))
    if res.get("tabla_lsd") is not None and not res["tabla_lsd"].empty:
        salida.append(("LSD", res["tabla_lsd"]))
        salida.append(("Letras LSD", res["letras_lsd"].rename(columns={"Media": f"Promedio ({u})"})))
    return salida



def _trozos_png(png: bytes, w, h, proporcion: float):
    """Corta una imagen alta en franjas cuya relación alto/ancho cabe en una hoja."""
    from PIL import Image as PILImage
    img = PILImage.open(io.BytesIO(png))
    W, H = img.size
    alto_trozo = max(1, int(W * proporcion))
    cortes = list(range(0, H, alto_trozo)) + [H]
    if len(cortes) > 2 and cortes[-1] - cortes[-2] < alto_trozo * 0.35:
        del cortes[-2]  # una franja final muy corta (solo el eje) se une a la anterior
    salida = []
    for y, y2 in zip(cortes[:-1], cortes[1:]):
        parte = img.crop((0, y, W, y2))
        buf = io.BytesIO()
        parte.save(buf, format="PNG")
        salida.append((buf.getvalue(), parte.size[0], parte.size[1]))
    return salida

def exportar_pdf(ruta: str, figuras: List[Tuple[str, object]], res: dict, proyecto: Optional[pd.DataFrame]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate
    import reporte_biogrow as REP

    fuente, negrita = REP._fuentes()
    est = {"titulo": ParagraphStyle("t", fontName=negrita, fontSize=16, spaceAfter=4),
           "sub": ParagraphStyle("s", fontName=fuente, fontSize=9.5, textColor=colors.HexColor("#475569"), spaceAfter=10),
           "h2": ParagraphStyle("h2", fontName=negrita, fontSize=12, textColor=colors.HexColor("#" + VERDE),
                                spaceBefore=8, spaceAfter=5),
           "normal": ParagraphStyle("n", fontName=fuente, fontSize=9.5, leading=13),
           "nota": ParagraphStyle("no", fontName=fuente, fontSize=8, leading=11)}
    pagina = landscape(letter)
    doc = SimpleDocTemplate(ruta, pagesize=pagina, leftMargin=30, rightMargin=30, topMargin=28, bottomMargin=30,
                            title="Gráficas BioGrow", author="BioGrow Analytics")
    ancho, alto_disp = pagina[0] - 60, pagina[1] - 58 - 30
    el = [Paragraph("BioGrow Analytics — Gráficas del análisis", est["titulo"]),
          Paragraph(REP._esc(f"{datetime.now():%Y-%m-%d %H:%M}  |  {res.get('nombre_var', '')} "
                             f"({res.get('u', '')})  |  Comparación: {res.get('agrupacion', '')}  |  "
                             f"{len(figuras)} gráfica(s)"), est["sub"])]
    if res and res.get("ok"):
        el.append(Paragraph(REP._esc(_dictamen(res)), est["normal"]))
    if proyecto is not None and not proyecto.empty:
        el.append(Paragraph("Información del proyecto", est["h2"]))
        el.append(REP._tabla(proyecto, est, fuente, negrita, ancho * 0.8, tam=8.5))
    for titulo, fig in figuras:
        png, w, h = _png(fig)
        ancho_img = ancho
        alto_img = ancho_img * h / w
        if alto_img <= alto_disp:
            el += [PageBreak(), Paragraph(REP._esc(titulo), est["h2"]),
                   Image(io.BytesIO(png), width=ancho_img, height=alto_img)]
        elif alto_img <= alto_disp * 1.4:  # un poco alta: se reduce para caber en una hoja
            el += [PageBreak(), Paragraph(REP._esc(titulo), est["h2"]),
                   Image(io.BytesIO(png), width=alto_disp * w / h, height=alto_disp)]
        else:  # muy alta (muchos pares): se reparte en varias hojas a lo ancho completo
            for k, trozo in enumerate(_trozos_png(png, w, h, alto_disp / ancho)):
                tw, th = trozo[1], trozo[2]
                sufijo = "" if k == 0 else f" (continuación {k + 1})"
                aw, ah = ancho, ancho * th / tw
                if ah > alto_disp:
                    aw, ah = alto_disp * tw / th, alto_disp
                el += [PageBreak(), Paragraph(REP._esc(titulo + sufijo), est["h2"]),
                       Image(io.BytesIO(trozo[0]), width=aw, height=ah)]
    doc.build(el)


def exportar_excel(ruta: str, figuras: List[Tuple[str, object]], res: dict, proyecto: Optional[pd.DataFrame],
                   incluir_tablas: bool = True) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Proyecto"
    ws["A1"] = "BioGrow Analytics — Gráficas y cálculos"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {datetime.now():%Y-%m-%d %H:%M}"
    if res and res.get("ok"):
        ws["A3"] = _dictamen(res)
        ws["A3"].font = Font(bold=True, color=VERDE if (res.get("anova") or {}).get("significativo") else "B91C1C")
    fila = 5
    if proyecto is not None and not proyecto.empty:
        for c, t in enumerate(["Dato", "Valor"], start=1):
            ws.cell(fila, c, t).font = Font(bold=True, color="FFFFFF")
            ws.cell(fila, c).fill = PatternFill("solid", fgColor=VERDE)
        for dato, valor in proyecto.itertuples(index=False):
            fila += 1
            ws.cell(fila, 1, dato).font = Font(bold=True)
            ws.cell(fila, 2, valor)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100

    wg = wb.create_sheet("Gráficas")
    fila = 1
    for titulo, fig in figuras:
        wg.cell(fila, 1, titulo).font = Font(bold=True, size=12, color=VERDE)
        png, w, h = _png(fig, dpi=110)
        img = XLImage(io.BytesIO(png))
        escala = min(1.0, 1100 / img.width)
        img.width, img.height = int(img.width * escala), int(img.height * escala)
        wg.add_image(img, f"A{fila + 1}")
        fila += int(img.height / 20) + 4

    if incluir_tablas:
        for nombre, tabla in tablas_calculos(res):
            hoja = wb.create_sheet(re.sub(r"[\\\\/*?:\\[\\]]", "-", nombre)[:31])
            cols, filas = AN.formatear_tabla(tabla)
            hoja.append(cols)
            for c in hoja[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=VERDE)
            for f in filas:
                hoja.append([_numero(v) for v in f])
            for i, c in enumerate(cols, start=1):
                largo = max([len(str(c))] + [len(f[i - 1]) for f in filas[:200]])
                hoja.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 50)
            hoja.freeze_panes = "A2"
        if res and res.get("ok"):
            hoja = wb.create_sheet("Datos de los grupos")
            datos = res["df_an"]
            cols = [c for c in ["Grupo", "Variedad", "Tratamiento", "Temperatura", "Gas", "Flujo", "Orientación",
                                "Morfología", "Bloque", "No", res["variable"]] if c in datos.columns]
            hoja.append([("Semilla n.º" if c == "No" else (f"{res['nombre_var']} ({res['u']})" if c == res["variable"] else c))
                         for c in cols])
            for c in hoja[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=VERDE)
            for r in datos[cols].itertuples(index=False):
                hoja.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
            for i in range(1, len(cols) + 1):
                hoja.column_dimensions[get_column_letter(i)].width = 22
            hoja.freeze_panes = "A2"
    wb.save(ruta)


def _numero(texto: str):
    try:
        return float(texto) if re.fullmatch(r"-?\d+(\.\d+)?", texto) else texto
    except ValueError:
        return texto

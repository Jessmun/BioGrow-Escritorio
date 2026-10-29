"""
motor_biogrow.py — Motor de lectura y análisis estadístico de BioGrow.

Copia fiel de las funciones de cálculo de dashboardst.py (dashboard de Streamlit), sin
Streamlit ni Plotly, para que la app de escritorio haga exactamente los mismos análisis:
lectura de Excel (todos los formatos), descartes, promedio por caja, descriptivos, ANOVA,
Tukey, LSD, letras y cuartiles. Generado automáticamente: si cambias un cálculo en
dashboardst.py, vuelve a generar este archivo en lugar de editarlo a mano.
Internamente las longitudes se leen en cm (igual que el dashboard).
"""
from __future__ import annotations

import colorsys

import hashlib

import io

import json

import math

import random

import re

import unicodedata

import zipfile

from datetime import datetime

from itertools import combinations

from pathlib import Path

from xml.etree import ElementTree as ET

import numpy as np

import pandas as pd

PATRON_TITULO = re.compile(r"^\s*T\s*[-–]\s*(0*)(\d+)\s*([MNmn]?)\s*$")

PATRON_FECHA_HOJA = re.compile(
    r"(\d{1,2})\s*(?:de\s+)?("
    r"enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
    r"septiembre|octubre|noviembre|diciembre)",
    re.IGNORECASE,
)

MES_A_NUMERO = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}

TRATAMIENTOS_POR_FECHA = {
    (28, 2): {"ambos": 120},           
    (7, 3): {"par": 60, "impar": 30}, 
    (14, 3): {"par": 60, "impar": 30},  
    (21, 3): {"par": 75, "impar": 45},  
    (25, 4): {"par": 75, "impar": 45},  
    (2, 5): {"par": 45, "impar": 30},   
    (16, 5): {"par": 45, "impar": 30},  # no está en la tabla del protocolo: confirmar
}

NS_MAIN = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

NS_REL = {
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

def _col_a_indice(letras: str) -> int:
    indice = 0
    for caracter in letras:
        indice = indice * 26 + (ord(caracter.upper()) - ord("A") + 1)
    return indice

def _partir_celda(referencia: str):
    coincidencia = re.match(r"([A-Z]+)(\d+)", referencia)
    letras, numero = coincidencia.groups()
    return letras, int(numero)

def _leer_cadenas_compartidas(zf: zipfile.ZipFile):
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []

    raiz = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    cadenas = []
    for si in raiz.findall("m:si", NS_MAIN):
        texto = "".join(t.text or "" for t in si.iter("{%s}t" % NS_MAIN["m"]))
        cadenas.append(texto)
    return cadenas

def _mapa_hojas(zf: zipfile.ZipFile):
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))

    rid_a_destino = {}
    for rel in rels.findall("pr:Relationship", NS_REL):
        rid_a_destino[rel.get("Id")] = rel.get("Target")

    hojas = []
    for hoja in workbook.find("m:sheets", NS_MAIN).findall("m:sheet", NS_MAIN):
        nombre = hoja.get("name")
        rid = hoja.get("{%s}id" % NS_REL["r"])
        destino = rid_a_destino.get(rid, "")
        if not destino.startswith("/"):
            destino = "xl/" + destino
        else:
            destino = destino.lstrip("/")
        hojas.append((nombre, destino))
    return hojas

def _leer_celdas(zf: zipfile.ZipFile, ruta_hoja: str, cadenas):
    raiz = ET.fromstring(zf.read(ruta_hoja))
    datos = raiz.find("m:sheetData", NS_MAIN)
    celdas = {}

    if datos is None:
        return celdas

    for fila in datos.findall("m:row", NS_MAIN):
        for celda in fila.findall("m:c", NS_MAIN):
            referencia = celda.get("r")
            if not referencia:
                continue

            letras, numero_fila = _partir_celda(referencia)
            col = _col_a_indice(letras)
            tipo = celda.get("t")
            v = celda.find("m:v", NS_MAIN)

            if v is None or v.text is None:
                is_ = celda.find("m:is", NS_MAIN)
                if is_ is not None:
                    texto = "".join(
                        t.text or "" for t in is_.iter("{%s}t" % NS_MAIN["m"])
                    )
                    celdas[(numero_fila, col)] = texto
                continue

            texto = v.text

            if tipo == "s":  
                celdas[(numero_fila, col)] = cadenas[int(texto)]
            elif tipo == "b":  
                celdas[(numero_fila, col)] = int(texto)
            else:
                try:
                    numero = float(texto)
                    celdas[(numero_fila, col)] = int(numero) if numero.is_integer() else numero
                except ValueError:
                    celdas[(numero_fila, col)] = texto

    return celdas

def _fecha_de_hoja(nombre_hoja: str):
    coincidencia = PATRON_FECHA_HOJA.search(nombre_hoja.lower())
    if not coincidencia:
        return None, None, None

    dia = int(coincidencia.group(1))
    mes_texto = coincidencia.group(2).lower()
    mes = MES_A_NUMERO[mes_texto]
    etiqueta = f"{dia:02d} {mes_texto}"
    return dia, mes, etiqueta

def _color_de_hoja(nombre_hoja: str) -> str:
    texto = nombre_hoja.lower()
    if "cacahuazintle" in texto:
        return "C"
    if "criollo" in texto:
        return "CR"
    if "rosa" in texto or re.search(r"datos-?r\b", texto) or "-r " in texto or "(r" in texto:
        return "R"
    if "blanc" in texto or re.search(r"datos-?b\b", texto):
        return "B"
    marca = re.search(r"datos-?([a-z]+)", texto)
    if marca:
        codigo = marca.group(1)
        if codigo in ("c", "cacahuazintle"):
            return "C"
        if codigo in ("cr", "criollo"):
            return "CR"
        if codigo == "r":
            return "R"
        if codigo == "b":
            return "B"
    return "?"

PATRON_LEYENDA = re.compile(r"^\s*(\d+)\s*[-–]\s*(\d+(?:[.,]\d+)?)\s*(s|seg|min)\b", re.IGNORECASE)

PATRON_TIEMPO = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*(s|seg|segundos|min|minutos)\s*$", re.IGNORECASE)

def _a_segundos(numero_txt, unidad):
    valor = float(str(numero_txt).replace(",", "."))
    return valor * 60 if unidad.lower().startswith("min") else valor

def _leyenda_paridad(celdas):
    """Lee la leyenda que algunas hojas traen en las primeras filas, p. ej.
    '1 - 45s' / '2 - 75s'  ->  {impar: 45, par: 75}.
    Tiene prioridad sobre la tabla fija TRATAMIENTOS_POR_FECHA, porque la tabla
    está indexada solo por fecha y hay fechas con dos variedades que recibieron
    tiempos distintos (p. ej. 2 de mayo: Blanca 30/45 s, Rosa 45/75 s)."""
    regla = {}
    for (fila, _col), valor in celdas.items():
        if fila > 2 or not isinstance(valor, str):
            continue
        m = PATRON_LEYENDA.match(valor.strip())
        if m:
            clave = "par" if int(m.group(1)) % 2 == 0 else "impar"
            regla[clave] = _a_segundos(m.group(2), m.group(3))
    return regla if {"par", "impar"} <= set(regla) else None

def _anotacion_bloque(celdas, fila_titulo, col_bloque):
    """Texto escrito justo encima de un bloque (fila anterior, en la columna del
    título o la siguiente). En los datos aparece como '2min', 'Sin aire' o
    'Plasma en el petri'. Devuelve (segundos_override, condición)."""
    for col in (col_bloque, col_bloque + 1):
        valor = celdas.get((fila_titulo - 1, col))
        if not isinstance(valor, str) or not valor.strip():
            continue
        texto = valor.strip()
        if PATRON_TITULO.match(texto) or PATRON_LEYENDA.match(texto):
            continue
        m = PATRON_TIEMPO.match(texto.replace(" ", ""))
        if m:
            return _a_segundos(m.group(1), m.group(2)), None
        return None, texto
    return None, None

def _segundos_tratamiento(dia, mes, tirada, regla_hoja=None):
    regla = regla_hoja or TRATAMIENTOS_POR_FECHA.get((dia, mes))

    if regla is None:
        return None  # sin información: mejor "Sin definir" que inventar un tiempo

    if "ambos" in regla:
        return regla["ambos"]
    return regla["par"] if tirada % 2 == 0 else regla["impar"]

def _datos_clase(letra: str):
    letra = (letra or "").upper()
    if letra == "M":
        return "M", "Embrión"
    if letra == "N":
        return "N", "Endospermo"
    return "S", "Patrón"

_OFFSETS = {"No": 0, "Raíz": 1, "Tallo": 2, "Germinó": 3}

def _extraer_bloques_de_hoja(nombre_hoja, celdas, advertencias):
    dia, mes, etiqueta_fecha = _fecha_de_hoja(nombre_hoja)
    color = _color_de_hoja(nombre_hoja)
    regla_hoja = _leyenda_paridad(celdas)
    regla_tabla = TRATAMIENTOS_POR_FECHA.get((dia, mes))
    if regla_hoja and regla_tabla and regla_hoja != regla_tabla:
        advertencias.append(
            f"{nombre_hoja}: la leyenda de la hoja indica impar={regla_hoja['impar']:g} s / "
            f"par={regla_hoja['par']:g} s y difiere de la tabla fija por fecha; "
            "se usó la leyenda de la hoja."
        )
    titulos = []
    for (fila, col), valor in celdas.items():
        if isinstance(valor, str):
            coincidencia = PATRON_TITULO.match(valor.strip())
            if coincidencia:
                cero = coincidencia.group(1)
                tirada = int(coincidencia.group(2))
                letra = coincidencia.group(3)
                titulos.append((fila, col, tirada, letra, cero, valor.strip()))

    titulos.sort()
    # Las corridas nuevas ya no separan embrión (M) y endospermo (N). Si la hoja
    # no tiene ningún bloque M/N, los bloques sin cero inicial (T-1, T-2…) son
    # semillas tratadas por el endospermo y solo T-01…T-05 son el control.
    hoja_con_mn = any(t[3] for t in titulos)
    if titulos and not hoja_con_mn and any(not t[4] for t in titulos):
        advertencias.append(
            f"{nombre_hoja}: la hoja no tiene bloques M/N. Los bloques T-1, T-2… se tomaron "
            "como semillas tratadas (endospermo) y T-01…T-05 como control."
        )

    registros = []
    contador_replica = {}

    for fila_titulo, col_bloque, tirada, letra, cero, titulo_texto in titulos:
        if letra:
            codigo_clase, nombre_clase = _datos_clase(letra)
        elif cero or hoja_con_mn:
            codigo_clase, nombre_clase = "S", "Patrón"
        else:
            codigo_clase, nombre_clase = "N", "Endospermo"
        seg_anot, condicion = _anotacion_bloque(celdas, fila_titulo, col_bloque)

        if codigo_clase == "S":
            segundos = 0
            fuente_tiempo = "Control: sin tratamiento (0 s)"
        elif seg_anot is not None:
            segundos = seg_anot
            condicion = condicion or f"Tiempo anotado ({seg_anot:g} s)"
            fuente_tiempo = f"Nota escrita sobre el bloque ({seg_anot:g} s)"
        else:
            segundos = _segundos_tratamiento(dia, mes, tirada, regla_hoja)
            paridad = "par" if tirada % 2 == 0 else "impar"
            if segundos is None:
                fuente_tiempo = "Sin definir: la fecha no está en las especificaciones"
            elif regla_hoja:
                fuente_tiempo = f"Leyenda de la hoja (tirada {paridad})"
            elif "ambos" in (TRATAMIENTOS_POR_FECHA.get((dia, mes)) or {}):
                fuente_tiempo = "Especificaciones por fecha (todas las tiradas)"
            elif (dia, mes) == (16, 5):
                fuente_tiempo = (f"Tabla del código original, tirada {paridad} "
                                 "(16 de mayo no está en las especificaciones)")
            else:
                fuente_tiempo = f"Especificaciones por fecha (tirada {paridad})"
        condicion = condicion or "Estándar"

        clave_replica = (tirada, codigo_clase)
        contador_replica[clave_replica] = contador_replica.get(clave_replica, 0) + 1
        numero_replica = contador_replica[clave_replica]

       
        fila_datos = fila_titulo + 2
        n_leidas = 0

        while True:
            celda_no = celdas.get((fila_datos, col_bloque + _OFFSETS["No"]))
            if not isinstance(celda_no, (int, float)):
                break

            raiz = celdas.get((fila_datos, col_bloque + _OFFSETS["Raíz"]))
            tallo = celdas.get((fila_datos, col_bloque + _OFFSETS["Tallo"]))
            germino = celdas.get((fila_datos, col_bloque + _OFFSETS["Germinó"]))

            registros.append(
                {
                    "Hoja": nombre_hoja,
                    "Fecha etiqueta": etiqueta_fecha or nombre_hoja,
                    "Día": dia,
                    "Mes": mes,
                    "Color": color,
                    "Tirada": tirada,
                    "Paridad": "Par" if tirada % 2 == 0 else "Impar",
                    "Clase": codigo_clase,
                    "Clase nombre": nombre_clase,
                    "Bloque": titulo_texto,
                    "Réplica": numero_replica,
                    "No": int(celda_no),
                    "Raíz": raiz if isinstance(raiz, (int, float)) else None,
                    "Tallo": tallo if isinstance(tallo, (int, float)) else None,
                    "Germinó": int(germino) if isinstance(germino, (int, float)) else None,
                    "Tratamiento (s)": segundos,
                    "Condición": condicion,
                    "Fuente del tiempo": fuente_tiempo,
                }
            )
            n_leidas += 1
            fila_datos += 1

        if n_leidas == 0:
            advertencias.append(
                f"{nombre_hoja}: el bloque '{titulo_texto}' no tenía filas de datos legibles."
            )

    if etiqueta_fecha is None:
        advertencias.append(
            f"{nombre_hoja}: no se reconoció una fecha en el nombre de la hoja; "
            "el tratamiento en segundos quedó sin definir para esta pestaña."
        )

    return registros

TIPO_TIEMPO = "Tiempo de tratamiento por fecha (Blanco / Rosa)"

TIPO_TIEMPO30 = "Tiempo 30 s por variedad (Cacahuazintle / Criollo)"

TIPO_TEMPERATURA = "Temperatura del tratamiento (Criollo)"

TIPO_COLUMNAS = "Temperatura: un tratamiento por columna (Raíz / Tallo)"

TIPO_LARGO = "Variedades: una semilla por fila (tratamiento y temperatura)"

LONG = "cm"

CONFIG_TIPO = {
    TIPO_TIEMPO: {
        "dims": ["Tratamiento", "Clase nombre", "Variedad", "Condición", "Tirada", "Bloque", "Hoja", "Archivo"],
        "default": ["Tratamiento"],
        "descripcion": (
            "Hojas 'Datos-B/R (fecha)'. **T**=tirada · **M**=embrión · **N**=endospermo · "
            "T-01…T-05 = **patrón sin tratamiento**. Tiempo por fecha y paridad de la tirada: "
            "28 feb = 120 s; 7 y 14 mar = impar 30 s / par 60 s; 21 mar y 25 abr = impar 45 s / "
            "par 75 s; 2 may = impar 30 s / par 45 s. Si la hoja trae su propia leyenda "
            "('1 - 45s / 2 - 75s') se usa esa. Bloques anotados ('2min', 'Sin aire', "
            "'Plasma en el petri') se separan en 'Condición'."
        ),
        "pregunta": "¿Qué tiempo de tratamiento favorece la germinación y el crecimiento?",
    },
    TIPO_TIEMPO30: {
        "dims": ["Clase nombre", "Variedad", "Tratamiento", "Tirada", "Bloque"],
        "default": ["Clase nombre"],
        "descripcion": (
            "Todas las semillas tratadas recibieron 30 s. Se compara la cara expuesta "
            "(**Embrión / Endospermo**) contra el **Control** y entre variedades. Los "
            "'Experimento 1, 2, 3' son las cajas (réplicas). Longitudes en el Excel en mm."
        ),
        "pregunta": "¿Importa qué cara de la semilla se expone? ¿Responde igual cada variedad?",
    },
    TIPO_TEMPERATURA: {
        "dims": ["Temperatura", "Bloque", "Hoja", "Archivo", "Tratamiento"],
        "default": ["Temperatura"],
        "descripcion": (
            "Nomenclatura **'Mx - Txx'**: x = número de muestra (caja), xx = temperatura del "
            "tratamiento (Tamb = ambiente). T-01…T-05 = **patrón sin tratamiento**. Las "
            "longitudes vienen en **cm**; si se mezclan con archivos en mm, todos se pasan a la misma unidad para que sean comparables "
            "con los demás archivos. Los bloques de plantilla vacíos se descartan."
        ),
        "pregunta": "¿Qué temperatura durante el tratamiento favorece la germinación y el crecimiento?",
    },
    TIPO_COLUMNAS: {
        "dims": ["Temperatura", "Bloque", "Archivo"],
        "default": ["Temperatura"],
        "descripcion": (
            "Hojas **'Raíz'** y **'Tallo'** con un tratamiento por columna (**T-01** = control sin "
            "tratamiento; **Mx - temperatura**, p. ej. M2-Amb, M4-35 °C) y una semilla por fila. "
            "Las longitudes vienen en **cm**. Como no hay cajas, cada dato es una semilla."
        ),
        "pregunta": "¿Qué temperatura del tratamiento da raíces o tallos más largos que el control (T-01)?",
    },
    TIPO_LARGO: {
        "dims": ["Combinación", "Tratamiento", "Variedad", "Temperatura", "Gas", "Flujo", "Orientación",
                 "Morfología", "Bloque", "Fecha etiqueta", "Archivo"],
        "default": ["Variedad", "Combinación"],
        "descripcion": (
            "Tabla con **una semilla por fila**: lote (caja), variedad, tratamiento, temperatura, raíz y "
            "tallo (en **mm** en el Excel), estado de germinación y fecha. "
            "**Control** = sin tratamiento. Cada tratamiento se nombra con su tiempo y temperatura "
            "(p. ej. '30 s · 35 °C'). Compara los tratamientos **dentro de cada variedad**: una lenteja "
            "no se compara con un maíz."
        ),
        "pregunta": "¿Qué tratamiento (tiempo y temperatura) favorece la germinación y el crecimiento de cada variedad?",
    },
}

PATRON_MUESTRA = re.compile(
    r"^\s*M\s*0*(\d+)\s*[-–]\s*T?\s*(amb(?:iente)?|\d+(?:[.,]\d+)?)\s*(?:°\s*C)?\s*$",
    re.IGNORECASE,
)

PATRON_FECHA_ARCHIVO = re.compile(r"(\d{1,2})[-_ ]([A-Za-zé]+)[-_ ](\d{4})")

PATRON_T_NOTA = re.compile(r"\bt\s*=\s*(\d+(?:[.,]\d+)?)\s*s\b", re.IGNORECASE)

PATRON_T_ARCHIVO = re.compile(r"(?:^|[_\-\s]|t)(\d+(?:[.,]\d+)?)\s*s(?:\b|_|\.|$)", re.IGNORECASE)

def _es_formato_temperatura(celdas):
    return any(isinstance(v, str) and PATRON_MUESTRA.match(v.strip()) for v in celdas.values())

def _unidad_longitud(celdas):
    """Detecta la nota 'Longitud dada en cm'. Los valores se leen tal cual (factor 1);
    la conversión de mm a cm se hace después, en _postprocesar."""
    for (fila, _c), valor in celdas.items():
        if fila <= 3 and isinstance(valor, str) and re.search(r"\bcm\b", valor, re.IGNORECASE):
            return 1.0, "cm"
    return 1.0, "mm"

def _tiempo_temperatura(celdas, nombre_archivo, advertencias):
    """Tiempo de tratamiento: nota en la hoja ('t=60s') > nombre del archivo ('45s')."""
    en_hoja = None
    for (fila, _c), valor in celdas.items():
        if fila <= 2 and isinstance(valor, str):
            m = PATRON_T_NOTA.search(valor)
            if m:
                en_hoja = float(m.group(1).replace(",", "."))
    m = PATRON_T_ARCHIVO.search(Path(nombre_archivo or "").stem)
    en_nombre = float(m.group(1).replace(",", ".")) if m else None

    if en_hoja is not None and en_nombre is not None and en_hoja != en_nombre:
        advertencias.append(
            f"El nombre del archivo dice {en_nombre:g} s pero la hoja indica t={en_hoja:g} s; "
            "se usó el valor de la hoja. Verifica cuál es el correcto."
        )
    if en_hoja is not None:
        return en_hoja, f"Nota de la hoja (t={en_hoja:g} s)"
    if en_nombre is not None:
        return en_nombre, f"Nombre del archivo ({en_nombre:g} s)"
    return None, "Sin definir"

def _etiqueta_desde_archivo(nombre_archivo, respaldo):
    m = PATRON_FECHA_ARCHIVO.search(Path(nombre_archivo or "").stem)
    if not m:
        return respaldo
    variedad = "Criollo" if "criollo" in (nombre_archivo or "").lower() else respaldo
    return f"{variedad} {int(m.group(1)):02d}-{m.group(2).capitalize()}-{m.group(3)}"

def _extraer_hoja_temperatura(nombre_hoja, celdas, advertencias, nombre_archivo=""):
    factor, unidad = _unidad_longitud(celdas)
    segundos, fuente_seg = _tiempo_temperatura(celdas, nombre_archivo, advertencias)
    color = _color_de_hoja(nombre_hoja)
    etiqueta_hoja = _etiqueta_desde_archivo(nombre_archivo, nombre_hoja)

    titulos = []
    for (fila, col), valor in celdas.items():
        if not isinstance(valor, str):
            continue
        texto = valor.strip()
        m = PATRON_MUESTRA.match(texto)
        if m:
            titulos.append((fila, col, "M", int(m.group(1)), m.group(2), texto))
            continue
        m = PATRON_TITULO.match(texto)
        if m:
            # T-01…T-05 = control; T-5M, T-7N… = bloques con letra
            extra = m.group(3) if (m.group(3) or not m.group(1)) else ""
            titulos.append((fila, col, "T", int(m.group(2)), extra, texto))
    titulos.sort()

    # columnas sueltas tipo 'R' / 'T' (sin No ni Germinó): no se pueden usar
    for (fila, col), valor in celdas.items():
        if (isinstance(valor, str) and valor.strip().upper() == "R"
                and str(celdas.get((fila, col + 1), "")).strip().upper() == "T"):
            n = sum(isinstance(celdas.get((f, col)), (int, float)) for f in range(fila + 1, fila + 40))
            advertencias.append(
                f"Hay {n} valores en columnas sueltas 'R'/'T' (celda {chr(64 + col)}{fila}) sin "
                "número de semilla ni 'Germinó' ni nombre de muestra; no se incluyeron."
            )

    registros, vacios, contador = [], [], {}
    for fila_t, col_t, tipo, numero, extra, texto in titulos:
        filas_bloque, fila = [], fila_t + 2
        while isinstance(celdas.get((fila, col_t)), (int, float)):
            raiz = celdas.get((fila, col_t + 1))
            tallo = celdas.get((fila, col_t + 2))
            germ = celdas.get((fila, col_t + 3))
            filas_bloque.append((
                int(celdas[(fila, col_t)]),
                raiz * factor if isinstance(raiz, (int, float)) else None,
                tallo * factor if isinstance(tallo, (int, float)) else None,
                int(germ) if isinstance(germ, (int, float)) else None,
            ))
            fila += 1

        if not filas_bloque:
            advertencias.append(f"{texto}: bloque sin filas de datos legibles.")
            continue
        if all(not (r or t or g) for _, r, t, g in filas_bloque):
            vacios.append(texto)  # plantilla sin usar
            continue

        if tipo == "M":
            temp_txt = extra.lower()
            if temp_txt.startswith("amb"):
                temp_num, temp_lbl = np.nan, "Ambiente"
            else:
                temp_num = float(temp_txt.replace(",", "."))
                temp_lbl = f"{temp_num:g} °C"
            codigo, nombre_clase = "T", "Tratada"
            seg = segundos
            tirada = numero
        else:
            codigo, nombre_clase = _datos_clase(extra)
            if codigo == "S":
                temp_num, temp_lbl, seg = np.nan, "Control", 0
            else:
                temp_num, temp_lbl, seg = np.nan, "Sin anotar", segundos
                advertencias.append(f"{texto}: bloque con datos pero sin temperatura anotada.")
            tirada = numero if codigo != "S" else 0

        clave = (tipo, texto.upper().replace(" ", ""))
        contador[clave] = contador.get(clave, 0) + 1

        for no, raiz, tallo, germ in filas_bloque:
            registros.append({
                "Hoja": etiqueta_hoja, "Fecha etiqueta": etiqueta_hoja,
                "Día": None, "Mes": None, "Color": color,
                "Tirada": tirada, "Paridad": "Par" if tirada % 2 == 0 else "Impar",
                "Clase": codigo, "Clase nombre": nombre_clase,
                "Bloque": re.sub(r"\s+", " ", texto), "Réplica": contador[clave],
                "No": no, "Raíz": raiz, "Tallo": tallo, "Germinó": germ,
                "Tratamiento (s)": seg, "Condición": "Estándar",
                "Temperatura": temp_lbl, "Temperatura (°C)": temp_num,
                "Unidad original": unidad,
                "Fuente del tiempo": ("Control: sin tratamiento (0 s)" if seg == 0 else fuente_seg),
            })

    if vacios:
        advertencias.append(
            f"Se descartaron {len(vacios)} bloques de plantilla vacíos (todo en 0): "
            + ", ".join(vacios)
        )
    if segundos is None:
        advertencias.append("No se encontró el tiempo de tratamiento ni en la hoja ni en el nombre.")
    return registros

def _variable_de_hoja(nombre_hoja):
    t = nombre_hoja.strip().lower()
    if re.search(r"ra[ií]z|radic", t):
        return "Raíz"
    if re.search(r"tallo|hipoc|vástago|vastago", t):
        return "Tallo"
    return None

def _es_formato_columnas(celdas):
    """Hoja 'ancha': la primera fila con texto trae un tratamiento por columna (T-01, M2-Amb,
    M3-22…) y debajo solo hay números, una semilla por fila. Sin columnas No/Germinó."""
    if not celdas:
        return False
    fila_enc = min(f for f, _ in celdas)
    encabezados = {c: v for (f, c), v in celdas.items() if f == fila_enc}
    if len(encabezados) < 2 or not all(isinstance(v, str) and v.strip() for v in encabezados.values()):
        return False
    resto = [(f, c, v) for (f, c), v in celdas.items() if f > fila_enc]
    if any(isinstance(v, str) and v.strip() for _f, _c, v in resto):
        return False
    numeros = [(c, v) for _f, c, v in resto if isinstance(v, (int, float))]
    return len(numeros) >= 5 and all(c in encabezados for c, _ in numeros)

def _tratamiento_de_columna(texto, segundos):
    """T-01 → control; 'M2-Amb' / 'M2-ambiente' → muestra 2 a temperatura ambiente;
    'M4-35°C' → muestra 4 a 35 °C. Devuelve (bloque, clase, nombre_clase, temp_lbl, temp_num, seg, tirada)."""
    m = PATRON_TITULO.match(texto)
    if m:
        return (f"T-{int(m.group(2)):02d}", "S", "Patrón", "Control", np.nan, 0, 0)
    m = PATRON_MUESTRA.match(texto)
    if m:
        n, temp = int(m.group(1)), m.group(2).lower()
        if temp.startswith("amb"):
            lbl, num, corto = "Ambiente", np.nan, "Amb"
        else:
            num = float(temp.replace(",", "."))
            lbl, corto = f"{num:g} °C", f"{num:g} °C"
        return (f"M{n} - {corto}", "T", "Tratada", lbl, num, segundos, n)
    limpio = re.sub(r"\s+", " ", texto.strip())
    if _es_control(limpio):
        return (limpio, "S", "Patrón", "Control", np.nan, 0, 0)
    return (limpio, "T", "Tratada", limpio, np.nan, segundos, 0)

def _extraer_hoja_columnas(nombre_hoja, celdas, advertencias, nombre_archivo=""):
    variable = _variable_de_hoja(nombre_hoja)
    if variable is None:
        advertencias.append(
            f"Hoja '{nombre_hoja}': tiene un tratamiento por columna, pero por su nombre no se sabe si "
            "mide raíz o tallo. Renómbrala 'Raíz' o 'Tallo' para que se analice.")
        return []
    nota_mm = any(isinstance(v, str) and re.search(r"\bmm\b", v, re.IGNORECASE)
                  for v in list(celdas.values()) + [nombre_hoja])
    unidad = "mm" if nota_mm else "cm"
    segundos, fuente_seg = _tiempo_temperatura(celdas, nombre_archivo, [])
    etiqueta_hoja = _etiqueta_desde_archivo(nombre_archivo, Path(nombre_archivo or nombre_hoja).stem)
    color = _color_de_hoja(nombre_archivo or "")
    fila_enc = min(f for f, _ in celdas)
    registros = []
    for (f, col), texto in sorted(celdas.items()):
        if f != fila_enc:
            continue
        bloque, codigo, nombre_clase, temp_lbl, temp_num, seg, tirada = _tratamiento_de_columna(texto, segundos)
        valores = [(fila, celdas[(fila, col)]) for fila in sorted(ff for ff, cc in celdas if cc == col and ff > fila_enc)
                   if isinstance(celdas[(fila, col)], (int, float))]
        if not valores:
            advertencias.append(f"Hoja '{nombre_hoja}', columna {texto}: sin datos.")
            continue
        for n, (_fila, valor) in enumerate(valores, start=1):
            registros.append({
                "Hoja": etiqueta_hoja, "Fecha etiqueta": etiqueta_hoja, "Día": None, "Mes": None,
                "Color": color, "Tirada": tirada, "Paridad": "Par" if tirada % 2 == 0 else "Impar",
                "Clase": codigo, "Clase nombre": nombre_clase, "Bloque": bloque, "Réplica": 1,
                "No": n, "Raíz": float(valor) if variable == "Raíz" else None,
                "Tallo": float(valor) if variable == "Tallo" else None, "Germinó": 1,
                "Tratamiento (s)": seg, "Condición": "Estándar",
                "Temperatura": temp_lbl, "Temperatura (°C)": temp_num, "Unidad original": unidad,
                "Fuente del tiempo": ("Control: sin tratamiento (0 s)" if seg == 0 else fuente_seg),
            })
    advertencias.append(
        f"Hoja '{nombre_hoja}' ({variable}): un tratamiento por columna y una semilla por fila; "
        f"se leyeron {len(registros)} semillas. Longitudes tomadas en {unidad}"
        + (" (sin nota de unidad en la hoja)" if unidad == "cm" else " y pasadas a cm") + ".")
    return registros

_COLUMNAS_LARGO = {
    "gas": r"\bgas\b|atm[oó]sfera",
    "flujo": r"flujo|caudal|l\s*/\s*min",
    "orientacion": r"orientaci|cara expuesta|posici[oó]n",
    "morfologia": r"morfolog|forma de (la )?semilla",
    "raiz": r"ra[ií]z|radícula|radicula",
    "tallo": r"tallo|hipoc[oó]tilo|v[aá]stago",
    "caja": r"lote|caja|corrida|bloque|placa|petri|unidad",
    "variedad": r"variedad|especie|cultivo",
    "no": r"muestra|^no\.?$|^n[º°o]\.?\s|semilla\s*n|^semilla$|^id$",
    "tratamiento": r"tratamiento|tiempo|exposici",
    "temperatura": r"temperatura|temp\b|°c",
    "estado": r"estado|germin",
    "fecha": r"fecha|d[ií]a",
    "replica": r"r[eé]plica|repetici",
}

def _mapear_columnas_largo(encabezados):
    mapa = {}
    for col, texto in sorted(encabezados.items()):
        t = str(texto).strip().lower()
        for clave, patron in _COLUMNAS_LARGO.items():
            if clave not in mapa and re.search(patron, t):
                mapa[clave] = col
                break
    return mapa

def _es_formato_largo(celdas):
    if not celdas:
        return False
    fila_enc = min(f for f, _ in celdas)
    enc = {c: v for (f, c), v in celdas.items() if f == fila_enc and isinstance(v, str) and v.strip()}
    if len(enc) < 4:
        return False
    m = _mapear_columnas_largo(enc)
    if not ({"raiz", "tallo"} <= set(m) and ({"caja", "tratamiento", "variedad"} & set(m))):
        return False
    numeros = sum(isinstance(celdas.get((f, m["raiz"])), (int, float))
                  for f in range(fila_enc + 1, fila_enc + 31))
    return numeros >= 10

def _unidad_columna(texto, valores):
    t = str(texto).lower()
    if re.search(r"\bmm\b", t):
        return "mm"
    if re.search(r"\bcm\b", t):
        return "cm"
    positivos = [v for v in valores if v > 0]
    return "mm" if positivos and float(np.median(positivos)) > 30 else "cm"

def _germino_de_texto(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v != 0)
    t = str(v or "").strip().lower()
    if not t:
        return None
    if t.startswith("no") or t in ("0", "f", "falso", "false", "n"):
        return 0
    if t.startswith(("germin", "s", "1", "v", "true")):
        return 1
    return None

def _segundos_de_texto(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    t = str(v or "").strip().lower()
    if not t or re.search(r"control|testigo|patr[oó]n|sin\s+trat", t):
        return 0.0 if t else None
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*(min|m\b)?", t)
    if not m:
        return None
    n = float(m.group(1).replace(",", "."))
    return n * 60 if m.group(2) else n

def _temperatura_de_texto(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{float(v):g} °C", float(v)
    t = str(v or "").strip()
    if not t:
        return "No registrada", np.nan
    if re.match(r"amb", t, re.IGNORECASE):
        return "Ambiente", np.nan
    m = re.search(r"\d+(?:[.,]\d+)?", t)
    if m:
        n = float(m.group().replace(",", "."))
        return f"{n:g} °C", n
    return t, np.nan

_GASES = {"aire": "Aire", "ox": "Oxígeno", "hel": "Helio", "arg": "Argón", "nit": "Nitrógeno",
          "co2": "CO₂", "dióx": "CO₂", "diox": "CO₂"}

def _gas_de_texto(v):
    t = str(v or "").strip()
    if not t:
        return "Sin dato"
    tl = t.lower()
    if re.match(r"(sin|ningun|no aplica|n/?a|—|-|control)", tl):
        return "Sin gas"
    for clave, nombre in _GASES.items():
        if tl.startswith(clave):
            return nombre
    return t[:1].upper() + t[1:]

def _flujo_de_texto(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    m = re.search(r"\d+(?:[.,]\d+)?", str(v or ""))
    return float(m.group().replace(",", ".")) if m else np.nan

def _orientacion_de_texto(v):
    t = str(v or "").strip().lower()
    if not t:
        return "Sin dato"
    if t.startswith("emb"):
        return "Embrión"
    if t.startswith("endo"):
        return "Endospermo"
    if t.startswith(("alea", "azar", "random")):
        return "Aleatorio"
    return str(v).strip().capitalize()

def _morfologia_de_texto(v):
    t = str(v or "").strip().lower()
    if not t:
        return "Sin dato"
    if t.startswith("plan"):
        return "Plana"
    if t.startswith(("bol", "redond", "esf")):
        return "Bola"
    return str(v).strip().capitalize()

def _nombre_variedad(v):
    t = re.sub(r"\s+", " ", str(v or "").strip())
    t = re.sub(r"^semillas?\s+de\s+", "", t, flags=re.IGNORECASE)
    return t[:1].upper() + t[1:] if t else "Sin especificar"

def _extraer_hoja_larga(nombre_hoja, celdas, advertencias, nombre_archivo=""):
    fila_enc = min(f for f, _ in celdas)
    enc = {c: v for (f, c), v in celdas.items() if f == fila_enc and isinstance(v, str)}
    m = _mapear_columnas_largo(enc)
    filas = sorted({f for f, _ in celdas if f > fila_enc})
    def valor(f, clave):
        return celdas.get((f, m[clave])) if clave in m else None
    num = lambda x: float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None
    u_raiz = _unidad_columna(enc[m["raiz"]], [num(valor(f, "raiz")) or 0 for f in filas])
    u_tallo = _unidad_columna(enc[m["tallo"]], [num(valor(f, "tallo")) or 0 for f in filas])
    if u_raiz != u_tallo:
        advertencias.append(f"Hoja '{nombre_hoja}': raíz en {u_raiz} y tallo en {u_tallo}; ambos se pasan a cm.")
    etiqueta_hoja = _etiqueta_desde_archivo(nombre_archivo, Path(nombre_archivo or nombre_hoja).stem)
    registros, contador_no = [], {}
    sin_estado = "estado" not in m
    for f in filas:
        raiz, tallo = num(valor(f, "raiz")), num(valor(f, "tallo"))
        if raiz is None and tallo is None and valor(f, "estado") is None:
            continue
        variedad = _nombre_variedad(valor(f, "variedad")) if "variedad" in m else "Sin especificar"
        seg = _segundos_de_texto(valor(f, "tratamiento"))
        temp_lbl, temp_num = _temperatura_de_texto(valor(f, "temperatura"))
        texto_trat = str(valor(f, "tratamiento") or "").strip()
        control = seg == 0 or bool(re.search(r"control|testigo|patr[oó]n", texto_trat, re.IGNORECASE))
        if control:
            seg = 0.0
        caja_txt = str(valor(f, "caja") or "").strip() or f"{variedad} · {texto_trat or 'sin tratamiento'} · {temp_lbl}"
        rep = valor(f, "replica")
        rep = int(rep) if isinstance(rep, (int, float)) else 1
        clave = (caja_txt, rep)
        contador_no[clave] = contador_no.get(clave, 0) + 1
        no = valor(f, "no")
        no = int(no) if isinstance(no, (int, float)) else contador_no[clave]
        germ = _germino_de_texto(valor(f, "estado")) if not sin_estado else None
        if raiz is not None and u_raiz == "cm":
            raiz = raiz * 10  # todo se guarda en mm y _postprocesar lo pasa a cm
        if tallo is not None and u_tallo == "cm":
            tallo = tallo * 10
        if control:
            etiqueta_trat = "Control (0 s)"
        elif seg is not None:
            etiqueta_trat = f"{seg:g} s · {temp_lbl}" if "temperatura" in m else f"{seg:g} s"
        else:
            etiqueta_trat = f"{texto_trat} · {temp_lbl}" if texto_trat else temp_lbl
        fecha = valor(f, "fecha")
        fecha_txt = str(fecha)[:10] if fecha not in (None, "") else ""
        flujo = _flujo_de_texto(valor(f, "flujo")) if "flujo" in m else np.nan
        registros.append({
            "Gas": _gas_de_texto(valor(f, "gas")) if "gas" in m else "Sin dato",
            "Flujo (L/min)": flujo,
            "Flujo": (f"{flujo:g} L/min" if pd.notna(flujo) else ("Sin flujo" if "flujo" in m else "Sin dato")),
            "Orientación": _orientacion_de_texto(valor(f, "orientacion")) if "orientacion" in m else "Sin dato",
            "Morfología": _morfologia_de_texto(valor(f, "morfologia")) if "morfologia" in m else "Sin dato",
            "Hoja": etiqueta_hoja, "Fecha etiqueta": fecha_txt or etiqueta_hoja, "Día": None, "Mes": None,
            "Color": variedad, "Tirada": 0 if control else 1, "Paridad": "Par" if control else "Impar",
            "Clase": "S" if control else "T", "Clase nombre": "Patrón" if control else "Tratada",
            "Bloque": re.sub(r"\s+", " ", caja_txt), "Réplica": rep, "No": no,
            "Raíz": raiz, "Tallo": tallo, "Germinó": germ,
            "Tratamiento (s)": seg, "Etiqueta tratamiento": etiqueta_trat, "Condición": "Estándar",
            "Temperatura": "Control" if control else temp_lbl, "Temperatura (°C)": np.nan if control else temp_num,
            "Temperatura registrada": temp_lbl,
            "Unidad original": "mm", "Fuente del tiempo": "Columna de tratamiento" if "tratamiento" in m else "Sin definir",
        })
    if not registros:
        advertencias.append(f"Hoja '{nombre_hoja}': no se encontraron semillas con datos.")
        return []
    faltan = [n for c, n in [("caja", "lote / caja"), ("tratamiento", "tratamiento"),
                             ("estado", "estado de germinación")] if c not in m]
    advertencias.append(
        f"Hoja '{nombre_hoja}': tabla con una semilla por fila; se leyeron {len(registros)} semillas "
        f"(raíz en {u_raiz}, tallo en {u_tallo})."
        + (f" No se encontró la columna de {', '.join(faltan)}." if faltan else ""))
    return registros

def _leer_libro(fuente):
    """Devuelve [(nombre_hoja, celdas)] desde una ruta o bytes de un .xlsx."""
    if isinstance(fuente, (bytes, bytearray)):
        fuente = io.BytesIO(fuente)
    hojas = []
    with zipfile.ZipFile(fuente) as zf:
        cadenas = _leer_cadenas_compartidas(zf)
        for nombre_hoja, ruta_hoja in _mapa_hojas(zf):
            hojas.append((nombre_hoja, _leer_celdas(zf, ruta_hoja, cadenas)))
    return hojas

def cargar_datos(fuente, nombre_archivo=None):
    """Lee un Excel (ruta o bytes), detecta el formato de cada hoja y extrae los datos."""
    if nombre_archivo is None and not isinstance(fuente, (bytes, bytearray)):
        nombre_archivo = Path(fuente).name
    advertencias, todos = [], []

    for nombre_hoja, celdas in _leer_libro(fuente):
        if not celdas:
            continue
        if re.match(r"\s*(l[eé]eme|instrucciones|readme|notas)\b", nombre_hoja, re.IGNORECASE):
            continue  # hojas de ayuda (p. ej. en la plantilla), no traen datos
        if _es_formato_largo(celdas):
            tipo = TIPO_LARGO
            registros = _extraer_hoja_larga(nombre_hoja, celdas, advertencias, nombre_archivo)
        elif _es_formato_columnas(celdas):
            tipo = TIPO_COLUMNAS
            registros = _extraer_hoja_columnas(nombre_hoja, celdas, advertencias, nombre_archivo)
        elif _es_formato_temperatura(celdas):
            tipo = TIPO_TEMPERATURA
            registros = _extraer_hoja_temperatura(nombre_hoja, celdas, advertencias, nombre_archivo)
        elif _es_formato_tiempo(celdas):
            tipo = TIPO_TIEMPO30
            registros = _extraer_hoja_tiempo(nombre_hoja, celdas, advertencias)
        else:
            tipo = TIPO_TIEMPO
            registros = _extraer_bloques_de_hoja(nombre_hoja, celdas, advertencias)
        if tipo in (TIPO_TIEMPO, TIPO_TIEMPO30) and _unidad_longitud(celdas)[1] == "cm":
            # hoja con la nota 'Longitud dada en cm' (p. ej. exportada desde este dashboard)
            for r in registros:
                r["Unidad original"] = "cm"
        for r in registros:
            r["Tipo de experimento"] = tipo
        todos.extend(registros)

    df = pd.DataFrame(todos)
    if df.empty:
        return df, advertencias
    return _postprocesar(df), advertencias

def _etiqueta_tratamiento(seg):
    if pd.isna(seg):
        return "Sin definir"
    if seg == 0:
        return "Control (0 s)"
    return f"{int(seg)} s" if float(seg).is_integer() else f"{seg:g} s"

def _postprocesar(df):
    for columna in ["Raíz", "Tallo", "Germinó", "Tratamiento (s)"]:
        df[columna] = pd.to_numeric(df[columna], errors="coerce")

    for col, defecto in [("Condición", "Estándar"), ("Temperatura", "No registrada"),
                         ("Fuente del tiempo", "Sin definir"),
                         ("Unidad original", "mm"), ("Tipo de experimento", TIPO_TIEMPO)]:
        if col not in df.columns:
            df[col] = defecto
        df[col] = df[col].fillna(defecto)
    if "Temperatura (°C)" not in df.columns:
        df["Temperatura (°C)"] = np.nan
    for col in ("Gas", "Flujo", "Orientación", "Morfología"):
        if col not in df.columns:
            df[col] = "Sin dato"
        df[col] = df[col].fillna("Sin dato")
    if "Flujo (L/min)" not in df.columns:
        df["Flujo (L/min)"] = np.nan
    # En los formatos con bloques M / N la orientación es la parte expuesta.
    sin_or = (df["Orientación"] == "Sin dato") & df["Clase nombre"].isin(["Embrión", "Endospermo"])
    df.loc[sin_or, "Orientación"] = df.loc[sin_or, "Clase nombre"]

    # Todas las longitudes se trabajan en cm: lo medido en cm se respeta tal cual
    # y lo medido en mm se divide entre 10.
    en_mm = df["Unidad original"] == "mm"
    df.loc[en_mm, ["Raíz", "Tallo"]] = (df.loc[en_mm, ["Raíz", "Tallo"]] / 10).round(4)

    # Coherencia: si hay longitud > 0 la semilla germinó.
    creció = (df["Raíz"].fillna(0) > 0) | (df["Tallo"].fillna(0) > 0)
    df.loc[df["Germinó"].isna() & creció, "Germinó"] = 1

    df["Tratamiento"] = df["Tratamiento (s)"].apply(_etiqueta_tratamiento)
    if "Etiqueta tratamiento" in df.columns:
        propia = df["Etiqueta tratamiento"].notna()
        df.loc[propia, "Tratamiento"] = df.loc[propia, "Etiqueta tratamiento"]
    df["Estado germinación"] = df["Germinó"].map({1: "Germinó", 0: "No germinó"})

    df["Corrida"] = (
        df["Color"].astype(str) + " | " + df["Bloque"].astype(str)
        + " | rep " + df["Réplica"].astype(str)
        + " | " + df["Hoja"].astype(str)
    )
    return df

FACTORES_COMBINACION = ["Temperatura", "Gas", "Flujo", "Orientación", "Morfología"]

def agregar_combinacion(df):
    """Columna 'Combinación': el tratamiento completo de cada lote (tiempo + los factores que
    cambian dentro de su variedad: temperatura, gas, flujo, orientación y morfología).
    Así, en la lenteja (sin variación de temperatura) se comparan combinaciones como
    '45 s · Argón · 2 L/min · Embrión'. Los factores constantes no se repiten en la etiqueta."""
    if df.empty:
        return df
    d = df.copy()
    tiempo = d["Tratamiento (s)"].apply(_etiqueta_tratamiento)
    control = d["Clase"].astype(str).eq("S") | tiempo.eq("Control (0 s)")
    vacios = {"Sin dato", "No registrada", "Sin anotar", "Control", "nan", ""}
    comb = pd.Series("", index=d.index, dtype=object)
    for variedad, idx in d.groupby(d["Variedad"].astype(str)).groups.items():
        sub = d.loc[idx]
        tratados = sub[~control.loc[idx]]
        usar = [c for c in FACTORES_COMBINACION if c in sub.columns
                and tratados[c].astype(str).loc[~tratados[c].astype(str).isin(vacios)].nunique() > 1]
        for i in idx:
            if control.loc[i]:
                partes = ["Control (0 s)"]
            else:
                partes = [tiempo.loc[i]] if tiempo.loc[i] != "Sin definir" else []
                partes += [str(d.at[i, c]) for c in usar if str(d.at[i, c]) not in vacios]
            comb.loc[i] = " · ".join(partes) or "Sin definir"
    d["Combinación"] = comb
    return d

def cargar_varios(fuentes):
    """fuentes: lista de rutas o de tuplas (nombre, bytes)."""
    marcos, advertencias = [], []

    for fuente in fuentes:
        if isinstance(fuente, tuple):
            nombre, contenido = fuente
        else:
            nombre, contenido = Path(fuente).name, fuente
        try:
            df_parcial, avisos = cargar_datos(contenido, nombre)
        except Exception as error:
            advertencias.append(f"No se pudo leer {nombre}: {error}")
            continue

        advertencias.extend(f"{nombre}: {a}" for a in avisos)
        if not df_parcial.empty:
            df_parcial = df_parcial.copy()
            df_parcial["Archivo"] = nombre
            marcos.append(df_parcial)
        else:
            advertencias.append(f"{nombre}: no se reconoció ningún bloque de datos.")

    if not marcos:
        return pd.DataFrame(), advertencias

    df = pd.concat(marcos, ignore_index=True)

    # Si dos archivos tienen hojas con el mismo nombre (p. ej. 'Datos-Criollo'),
    # se distinguen para que sus cajas no se mezclen.
    repetidas = df.groupby("Hoja")["Archivo"].nunique()
    repetidas = repetidas[repetidas > 1].index
    if len(repetidas):
        mask = df["Hoja"].isin(repetidas)
        df.loc[mask, "Hoja"] = df.loc[mask, "Hoja"] + " [" + df.loc[mask, "Archivo"].map(lambda x: Path(x).stem) + "]"
    df["Corrida"] = (
        df["Color"].astype(str) + " | " + df["Bloque"].astype(str)
        + " | rep " + df["Réplica"].astype(str) + " | " + df["Hoja"].astype(str)
    )
    # Identificador único de cada semilla (para descartes manuales)
    df["ID semilla"] = df["Archivo"].astype(str) + " | " + df["Corrida"] + " | No " + df["No"].astype(str)
    if df["ID semilla"].duplicated().any():
        df["ID semilla"] = df["ID semilla"] + " | fila " + df.groupby("ID semilla").cumcount().astype(str)
    return df, advertencias

def variedad_por_color(
    df,
    nombre_blanca="PBB-DH-183",
    nombre_rosa="H-392 A",
    nombre_cacahuazintle="CACAHUAZINTLE",
    nombre_criollo="CRIOLLO",
):
    mapa = {
        "B": nombre_blanca,
        "R": nombre_rosa,
        "C": nombre_cacahuazintle,
        "CR": nombre_criollo,
        "CACAHUAZINTLE": nombre_cacahuazintle,
        "CRIOLLO": nombre_criollo,
    }
    salida = df.copy()
    salida["Variedad"] = salida["Color"].apply(
        lambda valor: mapa.get(str(valor).strip().upper(), valor)
    )
    return salida

ENCABEZADO_RAIZ = "raíz [mm]"

TEXTO_SIN_GERMINAR = "sin germinar"

SEGUNDOS_TIEMPO30 = 30

def _clase_desde_texto(texto):
    t = (texto or "").strip().lower().rstrip(":")
    if t.startswith("embri"):
        return "M", "Embrión"
    if t.startswith("endosperm"):
        return "N", "Endospermo"
    return "S", "Patrón"

def _leer_bloques_tiempo30(celdas, col_datos, variedad, etiqueta_hoja, advertencias):
    registros = []

    encabezados = sorted(
        fila for (fila, col), valor in celdas.items()
        if col == col_datos and isinstance(valor, str)
        and valor.strip().lower() == ENCABEZADO_RAIZ
    )

    contador_replica = {}

    for fila_enc in encabezados:
        etiqueta_izq = celdas.get((fila_enc, col_datos - 1))
        etiqueta_izq = str(etiqueta_izq).strip() if etiqueta_izq else ""

        m = re.search(r"experimento\s*(\d+)", etiqueta_izq, re.IGNORECASE)
        if m:
            tirada = int(m.group(1))
            texto_clase = celdas.get((fila_enc - 1, col_datos)) or ""
            codigo, nombre_clase = _clase_desde_texto(str(texto_clase))
            nombre_bloque = f"Exp{tirada}-{codigo}"
        else:
            tirada = 0
            codigo, nombre_clase = "S", "Patrón"
            nombre_bloque = "Control"
        segundos_bloque = 0 if codigo == "S" else SEGUNDOS_TIEMPO30

        clave = (tirada, codigo)
        contador_replica[clave] = contador_replica.get(clave, 0) + 1
        replica = contador_replica[clave]
        fila = fila_enc + 1
        valores = []
        sin_germinar = 0
        limite = max(f for (f, _c) in celdas) + 2

        while fila < limite:
            celda = celdas.get((fila, col_datos))

            if isinstance(celda, str):
                if TEXTO_SIN_GERMINAR in celda.strip().lower():
                    conteo = celdas.get((fila, col_datos + 1))
                    sin_germinar = int(conteo) if isinstance(conteo, (int, float)) else 0
                break

            if isinstance(celda, (int, float)):
                tallo = celdas.get((fila, col_datos + 1))
                valores.append(
                    (float(celda), float(tallo) if isinstance(tallo, (int, float)) else None)
                )

            fila += 1

        if not valores and sin_germinar == 0:
            advertencias.append(
                f"{etiqueta_hoja} | {variedad} | {nombre_bloque}: bloque sin datos legibles."
            )
            continue

        numero = 0
        for raiz, tallo in valores:
            numero += 1
            registros.append({
                "Hoja": etiqueta_hoja,
                "Fecha etiqueta": etiqueta_hoja,
                "Día": None, "Mes": None,
                "Color": variedad,
                "Tirada": tirada,
                "Paridad": "Par" if tirada % 2 == 0 else "Impar",
                "Clase": codigo,
                "Clase nombre": nombre_clase,
                "Bloque": nombre_bloque,
                "Réplica": replica,
                "No": numero,
                "Raíz": raiz,
                "Tallo": tallo,
                "Germinó": 1,
                "Tratamiento (s)": segundos_bloque,
                "Condición": "Estándar",
                "Fuente del tiempo": ("Control: sin tratamiento (0 s)" if segundos_bloque == 0
                                      else "Encabezado del archivo (t = 30 s)"),
            })

        for _ in range(sin_germinar):
            numero += 1
            registros.append({
                "Hoja": etiqueta_hoja,
                "Fecha etiqueta": etiqueta_hoja,
                "Día": None, "Mes": None,
                "Color": variedad,
                "Tirada": tirada,
                "Paridad": "Par" if tirada % 2 == 0 else "Impar",
                "Clase": codigo,
                "Clase nombre": nombre_clase,
                "Bloque": nombre_bloque,
                "Réplica": replica,
                "No": numero,
                "Raíz": 0.0,   # se convierte a NaN en el análisis de crecimiento
                "Tallo": 0.0,
                "Germinó": 0,
                "Tratamiento (s)": segundos_bloque,
                "Condición": "Estándar",
                "Fuente del tiempo": ("Control: sin tratamiento (0 s)" if segundos_bloque == 0
                                      else "Encabezado del archivo (t = 30 s)"),
            })

    return registros

def _es_formato_tiempo(celdas):
    for valor in celdas.values():
        if isinstance(valor, str) and valor.strip().lower() == ENCABEZADO_RAIZ:
            return True
    return False

def _extraer_hoja_tiempo(nombre_hoja, celdas, advertencias):
    etiqueta = "Tiempo 30 s"
    for (fila, col), valor in celdas.items():
        if isinstance(valor, str) and re.match(r"^\s*t\s*=", valor.strip(), re.IGNORECASE):
            etiqueta = valor.strip().replace("t =", "Tiempo").strip()
            break

    variedades = []
    for (fila, col), valor in sorted(celdas.items()):
        if fila > 3 or not isinstance(valor, str):
            continue
        texto = valor.strip()
        if texto and texto.isupper() and not texto.lower().startswith("t ="):
            variedades.append((col, texto))

    if not variedades:
        advertencias.append(f"{nombre_hoja}: no se encontraron nombres de variedad.")
        return []

    registros = []
    for col_variedad, nombre_variedad in variedades:
        col_datos = col_variedad + 1
        encontrados = _leer_bloques_tiempo30(
            celdas, col_datos, nombre_variedad, etiqueta, advertencias
        )
        if not encontrados:
            encontrados = _leer_bloques_tiempo30(
                celdas, col_variedad, nombre_variedad, etiqueta, advertencias
            )
        registros.extend(encontrados)

    return registros

from scipy import stats

try:
    import statsmodels.formula.api as smf
    from statsmodels.stats.anova import anova_lm
    HAY_STATSMODELS = True
except Exception:  # la app sigue funcionando sin el ANOVA factorial
    HAY_STATSMODELS = False

SIG = "Diferencia significativa"

NO_SIG = "Sin diferencia significativa"

def ajustar_holm(p):
    """Corrección de Holm (controla el error familiar, más potente que Bonferroni)."""
    p = np.asarray(p, dtype=float)
    m = len(p)
    if m == 0:
        return p
    orden = np.argsort(p)
    ajust = np.empty(m)
    acumulado = 0.0
    for rango, idx in enumerate(orden):
        acumulado = max(acumulado, (m - rango) * p[idx])
        ajust[idx] = min(1.0, acumulado)
    return ajust

def _grupos_validos(df, variable, columna_grupo, minimo=2):
    grupos = []
    for nombre, sub in df.groupby(columna_grupo, dropna=False):
        valores = pd.to_numeric(sub[variable], errors="coerce").dropna()
        if len(valores) >= minimo:
            grupos.append((str(nombre), valores.astype(float)))
    return grupos

def preparar_variable(df, variable, incluir_no_germinadas=False):
    """Raíz y Tallo solo tienen sentido en semillas que germinaron.
    Una semilla no germinada no tiene longitud 0: no tiene longitud. Si se dejan
    los ceros, la media mezcla 'germinó o no' con 'cuánto creció' y la
    distribución queda inflada en cero (viola normalidad y homocedasticidad)."""
    salida = df.copy()
    if variable in ("Raíz", "Tallo") and not incluir_no_germinadas:
        salida.loc[salida["Germinó"] == 0, variable] = np.nan
    return salida

def agregar_por_unidad(df, variable, columnas_extra=()):
    """Pasa de semillas a unidades experimentales (una fila por caja/corrida).
    Las semillas de una misma caja comparten tratamiento, manipulación y
    microambiente: no son réplicas independientes. Tratar cada semilla como
    réplica infla los grados de libertad del error (pseudorreplicación)."""
    claves = ["Corrida", "Grupo"] + [c for c in columnas_extra if c not in ("Corrida", "Grupo")]
    claves = list(dict.fromkeys(claves))
    datos = df.copy()
    datos[variable] = pd.to_numeric(datos[variable], errors="coerce")
    agregado = (datos.groupby(claves, dropna=False)
                .agg(**{variable: (variable, "mean"),
                        "Semillas en la caja": (variable, "size"),
                        "Semillas con dato": (variable, "count")})
                .reset_index())
    agregado = agregado.dropna(subset=[variable])
    if variable == "Germinó":
        agregado[variable] = agregado[variable] * 100  # % de germinación por caja
    return agregado

def descriptivos_serie(serie, alfa=0.05):
    valores = pd.to_numeric(serie, errors="coerce").dropna()
    n = len(valores)

    base = {
        "n": n, "Media": np.nan, "Desviación estándar": np.nan,
        "Varianza": np.nan, "SEM": np.nan, "IC inf": np.nan, "IC sup": np.nan,
        "CV %": np.nan, "Q1": np.nan, "Mediana": np.nan,
        "Q3": np.nan, "IQR": np.nan, "Mínimo": np.nan, "Máximo": np.nan,
        "Atípicos (IQR)": 0,
    }
    if n == 0:
        return base

    media = valores.mean()
    sd = valores.std(ddof=1) if n > 1 else np.nan
    var = valores.var(ddof=1) if n > 1 else np.nan
    sem = sd / math.sqrt(n) if n > 1 else np.nan
    if n > 1:
        t = stats.t.ppf(1 - alfa / 2, n - 1)
        ic_inf, ic_sup = media - t * sem, media + t * sem
    else:
        ic_inf = ic_sup = np.nan

    q1 = valores.quantile(0.25)
    q3 = valores.quantile(0.75)
    iqr = q3 - q1
    atipicos = int(((valores < q1 - 1.5 * iqr) | (valores > q3 + 1.5 * iqr)).sum())

    base.update({
        "Media": media, "Desviación estándar": sd, "Varianza": var,
        "SEM": sem, "IC inf": ic_inf, "IC sup": ic_sup,
        "CV %": (sd / media * 100) if (n > 1 and media != 0) else np.nan,
        "Q1": q1, "Mediana": valores.median(), "Q3": q3,
        "IQR": iqr, "Mínimo": valores.min(), "Máximo": valores.max(),
        "Atípicos (IQR)": atipicos,
    })
    return base

def tabla_descriptivos(df, columna_grupo, variable, alfa=0.05):
    filas = []
    for nombre, sub in df.groupby(columna_grupo, dropna=False):
        fila = {"Grupo": str(nombre)}
        fila.update(descriptivos_serie(sub[variable], alfa))
        filas.append(fila)
    salida = pd.DataFrame(filas)
    if salida.empty:
        return salida
    return salida.sort_values("Media", ascending=False).reset_index(drop=True)

def ic_wilson(exitos, n, alfa=0.05):
    """IC de Wilson para una proporción (mejor que el de Wald con n pequeño o p cercana a 0/1)."""
    if n == 0:
        return np.nan, np.nan
    z = stats.norm.ppf(1 - alfa / 2)
    p = exitos / n
    den = 1 + z ** 2 / n
    centro = (p + z ** 2 / (2 * n)) / den
    medio = z * math.sqrt(p * (1 - p) / n + z ** 2 / (4 * n ** 2)) / den
    return max(0.0, centro - medio), min(1.0, centro + medio)

def supuestos(df, variable, columna_grupo, alfa=0.05):
    grupos = _grupos_validos(df, variable, columna_grupo)
    if len(grupos) < 2:
        return None

    residuos = np.concatenate([(v - v.mean()).values for _, v in grupos])
    n_res = len(residuos)
    if n_res < 3 or np.allclose(residuos, 0):
        return None

    if n_res <= 5000:
        w, p_norm = stats.shapiro(residuos)
        prueba_norm = "Shapiro-Wilk (residuos)"
    else:
        w, p_norm = stats.normaltest(residuos)
        prueba_norm = "D'Agostino-Pearson (residuos)"

    lev, p_lev = stats.levene(*[v.values for _, v in grupos], center="median")

    n_min = min(len(v) for _, v in grupos)
    normal = p_norm >= alfa
    homogeneo = p_lev >= alfa
    # Con ≥30 unidades por grupo el ANOVA es robusto a la no normalidad (TLC),
    # pero no a varianzas desiguales con n desbalanceado.
    normal_efectivo = normal or n_min >= 30

    if normal_efectivo and homogeneo:
        recomendado = "ANOVA + Tukey"
    elif normal_efectivo:
        recomendado = "Welch + Games-Howell"
    else:
        recomendado = "Kruskal-Wallis + Dunn"

    tabla = pd.DataFrame([
        {"Supuesto": "Normalidad", "Prueba": prueba_norm, "Estadístico": w,
         "p-valor": p_norm, "¿Se cumple?": "Sí" if normal else
         ("No (pero n≥30 por grupo: ANOVA robusto)" if n_min >= 30 else "No")},
        {"Supuesto": "Homogeneidad de varianzas", "Prueba": "Levene (Brown-Forsythe, mediana)",
         "Estadístico": lev, "p-valor": p_lev, "¿Se cumple?": "Sí" if homogeneo else "No"},
    ])
    return {"tabla": tabla, "residuos": residuos, "normal": normal,
            "homogeneo": homogeneo, "n mínimo": n_min, "recomendado": recomendado}

def anova_una_via(df, variable, columna_grupo, alfa=0.05):
    grupos = _grupos_validos(df, variable, columna_grupo)
    k = len(grupos)
    if k < 2:
        return None

    n_total = sum(len(v) for _, v in grupos)
    gran_media = np.concatenate([v.values for _, v in grupos]).mean()

    sc_entre = sum(len(v) * (v.mean() - gran_media) ** 2 for _, v in grupos)
    sc_dentro = sum(((v - v.mean()) ** 2).sum() for _, v in grupos)
    sc_total = sc_entre + sc_dentro

    gl_entre, gl_dentro, gl_total = k - 1, n_total - k, n_total - 1
    if gl_dentro <= 0:
        return None

    cm_entre = sc_entre / gl_entre
    cm_dentro = sc_dentro / gl_dentro
    f_calc = cm_entre / cm_dentro if cm_dentro > 0 else np.inf
    f_crit = stats.f.ppf(1 - alfa, gl_entre, gl_dentro)
    p_valor = stats.f.sf(f_calc, gl_entre, gl_dentro)
    significativo = p_valor < alfa

    eta2 = sc_entre / sc_total if sc_total > 0 else np.nan
    omega2 = ((sc_entre - gl_entre * cm_dentro) / (sc_total + cm_dentro)
              if sc_total > 0 else np.nan)

    col_fc = f"F crítica (α={alfa:g}; {gl_entre}, {gl_dentro})"
    signo = ">" if f_calc > f_crit else "≤"
    tabla = pd.DataFrame([
        {"Fuente": "Entre grupos (tratamiento)", "Suma de cuadrados": sc_entre,
         "GL": gl_entre, "Cuadrado medio": cm_entre, "F calculada": f_calc,
         col_fc: f_crit, "Comparación": f"{f_calc:.3f} {signo} {f_crit:.3f}",
         "p-valor": p_valor, "Resultado": SIG if significativo else NO_SIG},
        {"Fuente": "Dentro de grupos (error)", "Suma de cuadrados": sc_dentro,
         "GL": gl_dentro, "Cuadrado medio": cm_dentro, "F calculada": np.nan,
         col_fc: np.nan, "Comparación": "", "p-valor": np.nan, "Resultado": ""},
        {"Fuente": "Total", "Suma de cuadrados": sc_total, "GL": gl_total,
         "Cuadrado medio": np.nan, "F calculada": np.nan, col_fc: np.nan,
         "Comparación": "", "p-valor": np.nan, "Resultado": ""},
    ])

    medias = pd.DataFrame([
        {"Grupo": nombre, "n": len(v), "Media": v.mean(), "Varianza": v.var(ddof=1)}
        for nombre, v in grupos
    ])

    return {
        "metodo": "ANOVA de una vía", "tabla": tabla, "medias": medias,
        "CM error": cm_dentro, "GL error": gl_dentro, "GL entre": gl_entre,
        "k": k, "n total": n_total, "estadistico": f_calc,
        "nombre estadistico": "F", "F calculada": f_calc, "F crítica": f_crit,
        "p-valor": p_valor, "significativo": significativo, "alfa": alfa,
        "GL total": gl_total, "SC entre": sc_entre, "SC error": sc_dentro,
        "CM entre": cm_entre,
        "efecto": {"η² (eta cuadrada)": eta2, "ω² (omega cuadrada, menos sesgada)": omega2},
    }

def anova_welch(df, variable, columna_grupo, alfa=0.05):
    """ANOVA de Welch: no asume varianzas iguales."""
    grupos = _grupos_validos(df, variable, columna_grupo)
    k = len(grupos)
    if k < 2:
        return None
    n = np.array([len(v) for _, v in grupos], float)
    m = np.array([v.mean() for _, v in grupos])
    s2 = np.array([v.var(ddof=1) for _, v in grupos])
    if np.any(s2 <= 0):
        return None
    w = n / s2
    media_w = (w * m).sum() / w.sum()
    a = (w * (m - media_w) ** 2).sum() / (k - 1)
    lam = ((1 - w / w.sum()) ** 2 / (n - 1)).sum()
    b = 1 + 2 * (k - 2) / (k ** 2 - 1) * lam
    f = a / b
    gl1, gl2 = k - 1, (k ** 2 - 1) / (3 * lam)
    p = stats.f.sf(f, gl1, gl2)

    medias = pd.DataFrame([{"Grupo": g, "n": len(v), "Media": v.mean(), "Varianza": v.var(ddof=1)}
                           for g, v in grupos])
    todos = np.concatenate([v.values for _, v in grupos])
    sc_entre = sum(len(v) * (v.mean() - todos.mean()) ** 2 for _, v in grupos)
    sc_total = ((todos - todos.mean()) ** 2).sum()

    tabla = pd.DataFrame([{"Prueba": "ANOVA de Welch", "F": f, "GL num": gl1,
                           "GL den": gl2, "p-valor": p,
                           "Resultado": SIG if p < alfa else NO_SIG}])
    return {"metodo": "ANOVA de Welch", "tabla": tabla, "medias": medias,
            "k": k, "n total": int(n.sum()), "estadistico": f, "nombre estadistico": "F",
            "GL entre": gl1, "GL error": gl2, "p-valor": p, "significativo": p < alfa,
            "alfa": alfa,
            "efecto": {"η² (descriptiva)": sc_entre / sc_total if sc_total > 0 else np.nan}}

def kruskal_wallis(df, variable, columna_grupo, alfa=0.05):
    grupos = _grupos_validos(df, variable, columna_grupo)
    k = len(grupos)
    if k < 2:
        return None
    try:
        h, p = stats.kruskal(*[v.values for _, v in grupos])
    except ValueError:  # todos los valores idénticos
        return None
    n = sum(len(v) for _, v in grupos)
    eps2 = h / ((n ** 2 - 1) / (n + 1)) if n > 1 else np.nan
    medias = pd.DataFrame([{"Grupo": g, "n": len(v), "Media": v.mean(),
                            "Mediana": v.median()} for g, v in grupos])
    tabla = pd.DataFrame([{"Prueba": "Kruskal-Wallis", "H": h, "GL": k - 1,
                           "p-valor": p, "Resultado": SIG if p < alfa else NO_SIG}])
    return {"metodo": "Kruskal-Wallis", "tabla": tabla, "medias": medias, "k": k,
            "n total": n, "estadistico": h, "nombre estadistico": "H",
            "GL entre": k - 1, "GL error": np.nan, "p-valor": p,
            "significativo": p < alfa, "alfa": alfa,
            "efecto": {"ε² (épsilon cuadrada)": eps2}, "_grupos": grupos}

def tukey_hsd(anova):
    """Prueba de Tukey (Tukey-Kramer, válida con n desiguales).
    Para cada par se compara |diferencia de medias| contra la DHS
    (diferencia honestamente significativa) = q crítica × √(CMerror/2 · (1/n1 + 1/n2)).
    Es equivalente a comparar q calculada contra q crítica."""
    if anova is None or anova.get("metodo") != "ANOVA de una vía":
        return pd.DataFrame()
    medias, cm_error = anova["medias"], anova["CM error"]
    gl_error, k, alfa = anova["GL error"], anova["k"], anova["alfa"]
    if len(medias) < 2 or cm_error <= 0 or gl_error <= 0:
        return pd.DataFrame()

    q_crit = stats.studentized_range.ppf(1 - alfa, k, gl_error)
    filas = []
    for (_, g1), (_, g2) in combinations(medias.iterrows(), 2):
        dif = g1["Media"] - g2["Media"]
        error = math.sqrt((cm_error / 2) * (1 / g1["n"] + 1 / g2["n"]))
        q_calc = abs(dif) / error if error > 0 else np.inf
        dhs = q_crit * error
        p_aj = stats.studentized_range.sf(q_calc, k, gl_error)
        signif = abs(dif) > dhs
        filas.append({
            "Grupo 1": g1["Grupo"], "Grupo 2": g2["Grupo"],
            "n1": int(g1["n"]), "n2": int(g2["n"]),
            "Media 1": g1["Media"], "Media 2": g2["Media"],
            "|Diferencia|": abs(dif), "DHS": dhs,
            "Comparación |dif| vs DHS": f"{abs(dif):.3f} {'>' if signif else '≤'} {dhs:.3f}",
            "q calculada": q_calc, "q crítica": q_crit,
            "Comparación q": f"{q_calc:.3f} {'>' if signif else '≤'} {q_crit:.3f}",
            "k (medias)": k, "GL error": gl_error, "p ajustada": p_aj,
            "Resultado": SIG if signif else NO_SIG,
        })
    return pd.DataFrame(filas).sort_values("|Diferencia|", ascending=False).reset_index(drop=True)

def lsd(anova):
    """Prueba LSD de Fisher (diferencia mínima significativa).
    Para cada par se compara |diferencia de medias| contra
    LSD = t crítica(α/2; GL error) × √(CMerror · (1/n1 + 1/n2)),
    equivalente a comparar t calculada contra t crítica.
    Según Fisher, solo se interpreta si el ANOVA fue significativo; la tabla
    muestra la comparación siempre y el dashboard avisa cuando no aplica."""
    if anova is None or anova.get("metodo") != "ANOVA de una vía":
        return pd.DataFrame()
    medias, cm_error = anova["medias"], anova["CM error"]
    gl_error, alfa = anova["GL error"], anova["alfa"]
    if len(medias) < 2 or cm_error <= 0 or gl_error <= 0:
        return pd.DataFrame()

    t_crit = stats.t.ppf(1 - alfa / 2, gl_error)
    filas = []
    for (_, g1), (_, g2) in combinations(medias.iterrows(), 2):
        dif = abs(g1["Media"] - g2["Media"])
        error = math.sqrt(cm_error * (1 / g1["n"] + 1 / g2["n"]))
        t_calc = dif / error if error > 0 else np.inf
        valor_lsd = t_crit * error
        p = 2 * stats.t.sf(t_calc, gl_error)
        signif = dif > valor_lsd
        filas.append({
            "Grupo 1": g1["Grupo"], "Grupo 2": g2["Grupo"],
            "n1": int(g1["n"]), "n2": int(g2["n"]),
            "Media 1": g1["Media"], "Media 2": g2["Media"],
            "|Diferencia|": dif, "LSD": valor_lsd,
            "Comparación |dif| vs LSD": f"{dif:.3f} {'>' if signif else '≤'} {valor_lsd:.3f}",
            "t calculada": t_calc, "t crítica": t_crit,
            "Comparación t": f"{t_calc:.3f} {'>' if signif else '≤'} {t_crit:.3f}",
            "GL error": gl_error, "p-valor": p,
            "Resultado": SIG if signif else NO_SIG,
        })
    return pd.DataFrame(filas).sort_values("|Diferencia|", ascending=False).reset_index(drop=True)

def games_howell(resultado):
    """Post-hoc para varianzas desiguales (acompaña al ANOVA de Welch)."""
    if resultado is None:
        return pd.DataFrame()
    medias, k, alfa = resultado["medias"], resultado["k"], resultado["alfa"]
    filas = []
    for (_, g1), (_, g2) in combinations(medias.iterrows(), 2):
        v1, v2 = g1["Varianza"] / g1["n"], g2["Varianza"] / g2["n"]
        dif = g1["Media"] - g2["Media"]
        se = math.sqrt((v1 + v2) / 2)
        gl = (v1 + v2) ** 2 / (v1 ** 2 / (g1["n"] - 1) + v2 ** 2 / (g2["n"] - 1))
        q = abs(dif) / se if se > 0 else np.inf
        p = stats.studentized_range.sf(q, k, gl)
        medio_ic = stats.studentized_range.ppf(1 - alfa, k, gl) * se
        filas.append({"Grupo 1": g1["Grupo"], "Grupo 2": g2["Grupo"],
                      "n1": int(g1["n"]), "n2": int(g2["n"]),
                      "Media 1": g1["Media"], "Media 2": g2["Media"],
                      "Diferencia": abs(dif), "IC inf": dif - medio_ic, "IC sup": dif + medio_ic,
                      "q": q, "GL (Welch)": gl, "p ajustada": p,
                      "Resultado": SIG if p < alfa else NO_SIG})
    return pd.DataFrame(filas).sort_values("Diferencia", ascending=False).reset_index(drop=True)

def dunn(resultado):
    """Post-hoc de Dunn (rangos) con corrección de Holm; acompaña a Kruskal-Wallis."""
    if resultado is None:
        return pd.DataFrame()
    grupos, alfa = resultado["_grupos"], resultado["alfa"]
    todos = np.concatenate([v.values for _, v in grupos])
    rangos = stats.rankdata(todos)
    n = len(todos)
    _, conteos = np.unique(todos, return_counts=True)
    corr_empates = (conteos ** 3 - conteos).sum() / (12 * (n - 1)) if n > 1 else 0

    rango_medio, tam, inicio = {}, {}, 0
    for nombre, v in grupos:
        rango_medio[nombre] = rangos[inicio:inicio + len(v)].mean()
        tam[nombre] = len(v)
        inicio += len(v)

    filas = []
    for (a, _), (b, _) in combinations(grupos, 2):
        se = math.sqrt((n * (n + 1) / 12 - corr_empates) * (1 / tam[a] + 1 / tam[b]))
        z = (rango_medio[a] - rango_medio[b]) / se if se > 0 else 0.0
        filas.append({"Grupo 1": a, "Grupo 2": b, "n1": tam[a], "n2": tam[b],
                      "Rango medio 1": rango_medio[a], "Rango medio 2": rango_medio[b],
                      "z": abs(z), "p sin ajustar": 2 * stats.norm.sf(abs(z))})
    tabla = pd.DataFrame(filas)
    if tabla.empty:
        return tabla
    tabla["p ajustada (Holm)"] = ajustar_holm(tabla["p sin ajustar"].values)
    tabla["Resultado"] = np.where(tabla["p ajustada (Holm)"] < alfa, SIG, NO_SIG)
    return tabla.sort_values("z", ascending=False).reset_index(drop=True)

def chi2_germinacion(df, columna_grupo, alfa=0.05):
    datos = df.dropna(subset=["Germinó"])
    tabla = pd.crosstab(datos[columna_grupo], datos["Germinó"].astype(int))
    for c in (0, 1):
        if c not in tabla.columns:
            tabla[c] = 0
    tabla = tabla[[0, 1]]
    if len(tabla) < 2:
        return None

    resumen = []
    for g, fila in tabla.iterrows():
        n = int(fila.sum()); x = int(fila[1])
        li, ls = ic_wilson(x, n, alfa)
        resumen.append({"Grupo": str(g), "n": n, "Germinaron": x,
                        "Media": x / n * 100 if n else np.nan,  # % germinación
                        "IC inf %": li * 100, "IC sup %": ls * 100})
    resumen = pd.DataFrame(resumen)

    if (tabla[1].sum() == 0) or (tabla[0].sum() == 0):
        chi2, p, gl, esperadas = 0.0, 1.0, len(tabla) - 1, tabla.values.astype(float)
    else:
        chi2, p, gl, esperadas = stats.chi2_contingency(tabla.values, correction=False)
    n_tot = tabla.values.sum()
    v_cramer = math.sqrt(chi2 / n_tot) if n_tot else np.nan
    pct_bajas = float((np.asarray(esperadas) < 5).mean() * 100)

    salida = pd.DataFrame([{"Prueba": "χ² de independencia", "χ²": chi2, "GL": gl,
                            "p-valor": p, "V de Cramér": v_cramer,
                            "% celdas con esperado < 5": pct_bajas,
                            "Resultado": SIG if p < alfa else NO_SIG}])
    return {"metodo": "χ² de independencia", "tabla": salida, "medias": resumen,
            "k": len(tabla), "n total": int(n_tot), "estadistico": chi2,
            "nombre estadistico": "χ²", "GL entre": gl, "GL error": np.nan,
            "p-valor": p, "significativo": p < alfa, "alfa": alfa,
            "efecto": {"V de Cramér": v_cramer}, "_contingencia": tabla,
            "aviso esperadas": pct_bajas > 20}

def pares_proporciones(resultado_chi2):
    """Comparaciones por pares con prueba exacta de Fisher + corrección de Holm."""
    if resultado_chi2 is None:
        return pd.DataFrame()
    tabla, alfa = resultado_chi2["_contingencia"], resultado_chi2["alfa"]
    filas = []
    for a, b in combinations(tabla.index, 2):
        sub = tabla.loc[[a, b]].values
        _, p = stats.fisher_exact(sub)
        pa = sub[0, 1] / sub[0].sum() * 100
        pb = sub[1, 1] / sub[1].sum() * 100
        filas.append({"Grupo 1": str(a), "Grupo 2": str(b),
                      "% germ. 1": pa, "% germ. 2": pb,
                      "Diferencia (pp)": abs(pa - pb), "p sin ajustar": p})
    t = pd.DataFrame(filas)
    if t.empty:
        return t
    t["p ajustada (Holm)"] = ajustar_holm(t["p sin ajustar"].values)
    t["Resultado"] = np.where(t["p ajustada (Holm)"] < alfa, SIG, NO_SIG)
    return t.sort_values("Diferencia (pp)", ascending=False).reset_index(drop=True)

def anova_factorial(df, variable, factores, alfa=0.05):
    """ANOVA factorial tipo II con interacción. Unir varios factores en un solo
    'Grupo' y hacer ANOVA de una vía no permite separar efectos principales de
    la interacción (p. ej. ¿el efecto del tiempo depende de la cara expuesta?)."""
    if not HAY_STATSMODELS or len(factores) < 2:
        return None
    datos = df[[variable] + list(factores)].dropna().copy()
    if datos.empty:
        return None
    nombres = {}
    for i, f in enumerate(factores):
        nuevo = f"F{i}"
        datos[nuevo] = datos[f].astype(str)
        nombres[nuevo] = f
    datos = datos.rename(columns={variable: "Y"})
    niveles = [datos[c].nunique() for c in nombres]
    if min(niveles) < 2:
        return {"error": "Algún factor tiene un solo nivel con los filtros actuales."}

    celdas = datos.groupby(list(nombres)).size()
    esperadas = int(np.prod(niveles))
    incompleto = len(celdas) < esperadas
    termino = " * ".join(f"C({c})" for c in nombres)
    if incompleto:
        termino = " + ".join(f"C({c})" for c in nombres)  # sin interacción si faltan celdas
    import warnings
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            modelo = smf.ols(f"Y ~ {termino}", data=datos).fit()
            exog = modelo.model.exog
            if np.linalg.matrix_rank(exog) < exog.shape[1]:
                return {"error": (
                    "Los factores elegidos están confundidos (anidados): algún nivel de un "
                    "factor solo aparece junto con un nivel del otro. Ejemplo típico: "
                    "'Control (0 s)' solo existe en la clase 'Patrón'. Para el factorial "
                    "Tratamiento × Clase filtra Clase = Embrión y Endospermo (sin patrón).")}
            tabla = anova_lm(modelo, typ=2)
    except Exception as e:
        return {"error": f"No se pudo ajustar el modelo: {e}"}

    def renombrar(indice):
        texto = indice
        for corto, largo in nombres.items():
            texto = texto.replace(f"C({corto})", largo)
        return texto.replace("Residual", "Error").replace(":", " × ")

    tabla.index = [renombrar(i) for i in tabla.index]
    sc_error = tabla.loc["Error", "sum_sq"]
    tabla["η² parcial"] = tabla["sum_sq"] / (tabla["sum_sq"] + sc_error)
    tabla.loc["Error", "η² parcial"] = np.nan
    tabla["Resultado"] = np.where(tabla["PR(>F)"] < alfa, SIG, NO_SIG)
    tabla.loc["Error", "Resultado"] = ""
    tabla = tabla.rename(columns={"sum_sq": "Suma de cuadrados", "df": "GL",
                                  "F": "F calculada", "PR(>F)": "p-valor"})
    tabla = tabla.reset_index().rename(columns={"index": "Fuente"})
    return {"tabla": tabla, "incompleto": incompleto, "celdas": len(celdas),
            "esperadas": esperadas, "residuos": modelo.resid.values}

def letras_cld(medias, comparaciones_df):
    """Asigna letras de modo que:
      * dos grupos que NO difieren comparten al menos una letra, y
      * dos grupos que SÍ difieren no comparten ninguna.
    El algoritmo anterior (voraz) garantizaba lo segundo pero no lo primero:
    podía dejar sin letra común a pares que no diferían."""
    if medias is None or medias.empty or comparaciones_df is None or comparaciones_df.empty:
        return pd.DataFrame()

    medias = medias.sort_values("Media", ascending=False).reset_index(drop=True)
    grupos = medias["Grupo"].astype(str).tolist()

    distintos = set()
    for _, f in comparaciones_df.iterrows():
        if f["Resultado"] == SIG:
            distintos.add(frozenset((str(f["Grupo 1"]), str(f["Grupo 2"]))))

    columnas = [set(grupos)]
    for par in distintos:
        a, b = tuple(par)
        nuevas = []
        for col in columnas:
            if a in col and b in col:
                nuevas.append(col - {a})
                nuevas.append(col - {b})
            else:
                nuevas.append(col)
        # absorción: quitar columnas vacías, duplicadas o contenidas en otra
        nuevas = [c for c in nuevas if c]
        unicas = []
        for c in sorted(nuevas, key=len, reverse=True):
            if not any(c <= u for u in unicas):
                unicas.append(c)
        columnas = unicas

    # ordenar letras: 'a' para la columna que contiene la media más alta
    posicion = {g: i for i, g in enumerate(grupos)}
    columnas.sort(key=lambda c: min(posicion[g] for g in c))

    abecedario = "abcdefghijklmnopqrstuvwxyz"
    letras = {g: "" for g in grupos}
    for i, col in enumerate(columnas):
        letra = abecedario[i] if i < 26 else abecedario[i // 26 - 1] + abecedario[i % 26]
        for g in grupos:
            if g in col:
                letras[g] += letra

    salida = medias.copy()
    salida["Letras"] = salida["Grupo"].astype(str).map(letras)
    return salida[["Grupo", "n", "Media", "Letras"]]

def letras_tukey(anova, comparaciones_df):
    """Compatibilidad con el nombre anterior."""
    if anova is None:
        return pd.DataFrame()
    return letras_cld(anova["medias"], comparaciones_df)

def analisis_completo(df, variable, columna_grupo, metodo, alfa=0.05, nivel_semilla_binario=False):
    """Ejecuta prueba global + post-hoc según el método elegido.
    Devuelve dict con: global, posthoc, nombre_posthoc, letras, supuestos, metodo_usado."""
    if nivel_semilla_binario:
        glob = chi2_germinacion(df, columna_grupo, alfa)
        post = pares_proporciones(glob)
        return {"global": glob, "posthoc": post, "nombre_posthoc": "Fisher exacta por pares (Holm)",
                "letras": letras_cld(glob["medias"], post) if glob else pd.DataFrame(),
                "supuestos": None, "metodo_usado": "χ² + Fisher (Holm)"}

    sup = supuestos(df, variable, columna_grupo, alfa)
    usado = sup["recomendado"] if (metodo.startswith("Automático") and sup) else metodo
    if usado.startswith("Automático"):
        usado = "ANOVA + Tukey"

    if usado == "Welch + Games-Howell":
        glob = anova_welch(df, variable, columna_grupo, alfa)
        post, nombre = games_howell(glob), "Games-Howell"
    elif usado == "Kruskal-Wallis + Dunn":
        glob = kruskal_wallis(df, variable, columna_grupo, alfa)
        post, nombre = dunn(glob), "Dunn (Holm)"
    else:
        glob = anova_una_via(df, variable, columna_grupo, alfa)
        post, nombre = tukey_hsd(glob), "Tukey HSD"
        usado = "ANOVA + Tukey"

    return {"global": glob, "posthoc": post, "nombre_posthoc": nombre,
            "letras": letras_cld(glob["medias"], post) if glob else pd.DataFrame(),
            "supuestos": sup, "metodo_usado": usado}

PALETA_BASE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]

PALETA = list(PALETA_BASE)

COLOR_PRINCIPAL = PALETA[0]

COLOR_CONTROL = "#8a8984"

def generar_paleta(semilla, n=16):
    rng = random.Random(semilla)
    inicio = rng.random()
    saturacion = rng.uniform(0.6, 0.8)
    luces = (rng.uniform(0.42, 0.48), rng.uniform(0.58, 0.64))
    orden = [0, 3, 6, 1, 4, 7, 2, 5]
    colores = []
    for i in range(n):
        h = (inicio + orden[i % 8] / 8 + (i // 8) / 16) % 1.0
        r, g, b = colorsys.hls_to_rgb(h, luces[(i + i // 8) % 2], saturacion)
        colores.append("#%02x%02x%02x" % (round(r * 255), round(g * 255), round(b * 255)))
    return colores

def establecer_colores(semilla=None):
    global PALETA, COLOR_PRINCIPAL
    PALETA = list(PALETA_BASE) if semilla is None else generar_paleta(semilla)
    COLOR_PRINCIPAL = PALETA[0]

def paleta_aleatoria(n, semilla=0):
    return [PALETA[i % len(PALETA)] for i in range(max(n, 1))]

def obtener_paleta(nombre_o_semilla=0, n=12):
    try:
        semilla = int(nombre_o_semilla)
    except (TypeError, ValueError):
        semilla = 0
    return paleta_aleatoria(max(n, 12), semilla)

def _dispersion_columnas(resumen, dispersion):
    if dispersion == "Cuartiles (IQR)":
        arriba = (resumen["Q3"] - resumen["Media"]).clip(lower=0)
        abajo = (resumen["Media"] - resumen["Q1"]).clip(lower=0)
        return arriba, abajo, "cuartiles (Q1–Q3)", None
    if dispersion == "Error estándar (SEM)":
        return resumen["SEM"], None, "error estándar (SEM)", "SEM"
    if dispersion == "IC 95% de la media":
        arriba = (resumen["IC sup"] - resumen["Media"]).clip(lower=0)
        abajo = (resumen["Media"] - resumen["IC inf"]).clip(lower=0)
        return arriba, abajo, "IC 95% de la media", None
    return resumen["Desviación estándar"], None, "desviación estándar", "Desviación estándar"

def describir_prueba(res):
    if res is None:
        return ""
    est = res.get("estadistico", np.nan)
    gl1, gl2 = res.get("GL entre"), res.get("GL error")
    gl_txt = f"GL={gl1:g}" if gl1 is not None and pd.notna(gl1) else ""
    if gl2 is not None and pd.notna(gl2):
        gl_txt = f"GL=({gl1:g}, {gl2:.4g})"
    return (f"{res['metodo']}: {res['nombre estadistico']}={est:.3f} · {gl_txt} · "
            f"p={res['p-valor']:.4g}")

CARPETA = Path(__file__).resolve().parent

def encontrar_excels():
    return sorted(p for p in CARPETA.glob("*.xlsx") if not p.name.startswith("~$"))

def redondear(df, decimales=4):
    salida = df.copy()
    for c in salida.select_dtypes(include="number").columns:
        salida[c] = salida[c].round(decimales)
    return salida

def construir_grupo(df, dimensiones):
    if not dimensiones:
        dimensiones = ["Bloque"]

  
    prefijos = {
        "Tirada": "Tirada ",
        "Réplica": "Réplica ",
        "No": "Semilla ",
    }

    partes = []
    for dim in dimensiones:
        if dim == "Tirada":
            serie = df[dim].apply(
                lambda v: "Control (patrón)" if str(v) in ("0", "0.0") else f"Tirada {v}"
            )
        elif dim in prefijos:
            serie = prefijos[dim] + df[dim].astype(str)
        else:
            serie = df[dim].astype(str)
        partes.append(serie)

    if len(partes) == 1:
        return partes[0]

    salida = partes[0]
    for extra in partes[1:]:
        salida = salida + " | " + extra
    return salida

COLUMNAS_CAJA = ["Hoja", "Archivo", "Variedad", "Clase nombre", "Tratamiento",
                 "Tratamiento (s)", "Condición", "Tirada", "Bloque", "Réplica", "Color",
                 "Temperatura", "Tipo de experimento"]

COLUMNAS_REGISTRO = ["ID semilla", "Hoja", "Bloque", "Réplica", "No", "Grupo",
                     "Germinó", "Raíz", "Tallo"]

def cajas_vacias(df):
    """Cajas en las que ninguna semilla tiene longitud registrada (bloques de
    plantilla que no se realizaron, capturados con Germinó = 0)."""
    largo = (pd.to_numeric(df["Raíz"], errors="coerce").fillna(0) > 0) | \
            (pd.to_numeric(df["Tallo"], errors="coerce").fillna(0) > 0)
    con_dato = largo.groupby(df["Corrida"]).any()
    return set(con_dato[~con_dato].index)

def aplicar_descartes(df, variable, op):
    """Aplica los descartes elegidos y devuelve (datos, registro).

    Hay dos tipos de descarte:
      * Se quita la semilla completa (caja vacía, caja o semilla excluida por el
        usuario): deja de contar también en el % de germinación.
      * Se quita solo el valor de la variable (no germinó, cero, fuera de rango,
        atípico): la semilla sigue contando para el % de germinación.
    """
    d = df.copy()
    d[variable] = pd.to_numeric(d[variable], errors="coerce")
    registro = []

    def anotar(mascara, motivo, quita_semilla):
        if mascara.any():
            r = d.loc[mascara, [c for c in COLUMNAS_REGISTRO if c in d.columns]].copy()
            r["Motivo"] = motivo if isinstance(motivo, str) else motivo[mascara].values
            r["Efecto"] = ("Semilla excluida del análisis" if quita_semilla
                           else f"Valor de {variable} excluido")
            registro.append(r)

    # --- se quita la semilla completa -------------------------------------
    if op.get("cajas_vacias"):
        m = d["Corrida"].isin(cajas_vacias(d))
        anotar(m, "Caja vacía (sin ninguna longitud registrada)", True)
        d = d[~m]
    if op.get("cajas_excluidas"):
        m = d["Corrida"].isin(op["cajas_excluidas"])
        anotar(m, "Caja excluida por el usuario", True)
        d = d[~m]
    if op.get("numeros_excluidos"):
        m = d["No"].isin(op["numeros_excluidos"])
        anotar(m, "Número de semilla excluido por el usuario", True)
        d = d[~m]
    manuales = op.get("manuales") or {}
    if manuales:
        m = d["ID semilla"].isin(manuales.keys())
        anotar(m, d["ID semilla"].map(manuales).fillna("Descarte manual"), True)
        d = d[~m]

    # --- se quita solo el valor de la variable ----------------------------
    if variable in ("Raíz", "Tallo"):
        no_germ = d["Germinó"] == 0
        if op.get("no_germinadas", "excluir") == "excluir":
            anotar(no_germ & d[variable].notna(), "No germinó", False)
            d.loc[no_germ, variable] = np.nan
        else:
            d.loc[no_germ, variable] = d.loc[no_germ, variable].fillna(0)

        if op.get("ceros_germinadas"):
            m = (d["Germinó"] == 1) & (d[variable] == 0)
            anotar(m, f"Germinó pero {variable} = 0", False)
            d.loc[m, variable] = np.nan
            if variable == "Raíz":
                d["Germinó (conteo)"] = d["Germinó"].where(~m, 0)

        rango = op.get("rango")
        if rango:
            m = d[variable].notna() & ((d[variable] < rango[0]) | (d[variable] > rango[1]))
            anotar(m, f"Fuera del rango {rango[0]:g}–{rango[1]:g} {LONG}", False)
            d.loc[m, variable] = np.nan

        if op.get("atipicos"):
            factor = op.get("factor_iqr", 1.5)
            q1 = d.groupby("Grupo")[variable].transform(lambda x: x.quantile(0.25))
            q3 = d.groupby("Grupo")[variable].transform(lambda x: x.quantile(0.75))
            iqr = q3 - q1
            m = d[variable].notna() & ((d[variable] < q1 - factor * iqr) |
                                       (d[variable] > q3 + factor * iqr))
            anotar(m, f"Atípico en su grupo (regla {factor:g}×IQR)", False)
            d.loc[m, variable] = np.nan

    reg = pd.concat(registro, ignore_index=True) if registro else pd.DataFrame(
        columns=COLUMNAS_REGISTRO + ["Motivo", "Efecto"])
    return d, reg

def tabla_correcciones(op, variable, hojas_cm, manuales, registro, df_semilla, advertencias):
    def cuenta(prefijo):
        if registro.empty:
            return 0
        return int(registro["Motivo"].astype(str).str.startswith(prefijo).sum())

    filas = []
    plantillas = [a for a in advertencias if "bloques de plantilla vacíos" in a]
    n_pl = sum(len(a.split("(todo en 0):")[-1].split(",")) for a in plantillas) if plantillas else 0
    filas.append(("Bloques de plantilla vacíos (todo en 0)", "Sí, siempre",
                  "Al leer el Excel no se toman los bloques sin usar.",
                  f"{n_pl} bloque(s)" if n_pl else "Ninguno"))
    sueltas = [a for a in advertencias if "columnas sueltas" in a]
    if sueltas:
        n_s = sum(int(re.search(r"Hay (\d+)", a).group(1)) for a in sueltas if re.search(r"Hay (\d+)", a))
        filas.append(("Columnas sueltas 'R'/'T' sin número ni 'Germinó'", "Sí, siempre",
                      "No se pueden asignar a una caja; no se incluyen.", f"{n_s} valor(es)"))
    filas.append(("Hojas que ya estaban en cm sin indicarlo (no dividir entre 10)", "Sí: " + ", ".join(hojas_cm) if hojas_cm else "No",
                  "Multiplica por 10 las longitudes de esas hojas.", "—"))
    n = cuenta("Caja vacía")
    filas.append(("Excluir cajas vacías (no realizadas)", "Sí" if op.get("cajas_vacias") else "No",
                  "Quita cajas sin ninguna longitud registrada (no cuentan en la germinación).",
                  f"{n} semilla(s)"))
    if variable in ("Raíz", "Tallo"):
        v = variable.lower()
        if op.get("no_germinadas") == "excluir":
            filas.append(("Semillas que no germinaron", "Excluidas del promedio",
                          f"No entran al promedio de {v}; sí cuentan en el % de germinación.",
                          f"{cuenta('No germinó')} semilla(s)"))
        else:
            n_ng = int((df_semilla["Germinó"] == 0).sum())
            filas.append(("Semillas que no germinaron", "Incluidas con longitud 0",
                          f"Entran al promedio de {v} con 0 {LONG}.", f"{n_ng} semilla(s)"))
        filas.append((f"Excluir germinadas con {v} = 0", "Sí" if op.get("ceros_germinadas") else "No",
                      f"Germinadas con {v} = 0 salen del promedio"
                      + ("; en el % de germinación cuentan como no germinadas." if variable == "Raíz" else "."),
                      f"{cuenta('Germinó pero')} semilla(s)"))
        r = op.get("rango")
        filas.append(("Excluir valores fuera de un rango", f"Sí: {r[0]:g}–{r[1]:g} {LONG}" if r else "No",
                      f"Valores de {v} fuera del rango salen del promedio.", f"{cuenta('Fuera del rango')} valor(es)"))
        filas.append(("Excluir valores atípicos de cada grupo",
                      f"Sí: regla {op.get('factor_iqr', 1.5):g} × IQR" if op.get("atipicos") else "No",
                      f"Valores de {v} fuera de Q1 − k·IQR y Q3 + k·IQR de su grupo salen del promedio.",
                      f"{cuenta('Atípico')} valor(es)"))
    n_m = len(registro) - sum(cuenta(p) for p in ("Caja vacía", "Caja excluida", "Número de semilla",
                                                   "No germinó", "Germinó pero", "Fuera del rango", "Atípico"))
    filas.append(("Semillas descartadas a mano", f"Sí: {len(manuales)}" if manuales else "No",
                  "Semillas quitadas por el usuario en la pestaña Descartes.", f"{max(n_m, 0)} semilla(s)"))
    return pd.DataFrame(filas, columns=["Corrección", "¿Aplicada?", "Qué hace", "Datos afectados"])

def revisar_calidad(df):
    """Inconsistencias en los datos crudos que conviene corregir en el Excel."""
    avisos = []
    malas = df[(df["Germinó"] == 0) & ((df["Raíz"].fillna(0) > 0) | (df["Tallo"].fillna(0) > 0))]
    for _, f in malas.iterrows():
        avisos.append(f"{f['Hoja']} · {f['Bloque']} · semilla {f['No']}: Germinó=0 pero tiene "
                      f"Raíz={f['Raíz']} / Tallo={f['Tallo']}.")
    ceros = df[(df["Germinó"] == 1) & (df["Raíz"].fillna(0) == 0) & (df["Tallo"].fillna(0) == 0)]
    for _, f in ceros.iterrows():
        avisos.append(f"{f['Hoja']} · {f['Bloque']} · semilla {f['No']}: Germinó=1 pero Raíz y "
                      "Tallo son 0. Verifica el registro.")
    # cajas con exactamente los mismos valores (posible copia)
    for hoja, sub in df.groupby("Hoja"):
        firmas = {}
        for corrida, c in sub.groupby("Corrida"):
            vals = tuple(zip(c["Raíz"].fillna(-1).round(3), c["Tallo"].fillna(-1).round(3)))
            if any(v[0] > 0 or v[1] > 0 for v in vals):
                firmas.setdefault(vals, []).append(corrida)
        for grupo in firmas.values():
            if len(grupo) > 1:
                avisos.append(f"{hoja}: las cajas {' y '.join(g.split(' | ')[1] for g in grupo)} "
                              "tienen exactamente los mismos valores (posible copia). "
                              "Puedes excluir una en la pestaña '🧹 Descartes'.")
    # hojas que parecen medidas en cm sin indicarlo
    for hoja, sub in df[df["Unidad original"] == "mm"].groupby("Hoja"):
        g = sub[sub["Germinó"] == 1]
        if len(g) > 20 and g["Raíz"].quantile(0.9) < 3 and g["Tallo"].quantile(0.9) < 3:
            avisos.append(f"{hoja}: las longitudes parecen estar en cm (90 % de las raíces "
                          f"< {g['Raíz'].quantile(0.9) * 10:.1f} en la hoja) y la hoja no lo indica. Si es así, "
                          "márcala en '5 · Descartes y correcciones → Hojas que ya estaban en cm'.")
    return avisos

def hojas_probables_cm(df):
    salida = []
    for hoja, sub in df[df["Unidad original"] == "mm"].groupby("Hoja"):
        g = sub[sub["Germinó"] == 1]
        if len(g) > 20 and g["Raíz"].quantile(0.9) < 3 and g["Tallo"].quantile(0.9) < 3:
            salida.append(hoja)
    return salida

def pruebas_requeridas(datos, variable, alfa):
    """ANOVA de una vía + Tukey + LSD + letras de ambas pruebas."""
    anova = anova_una_via(datos, variable, "Grupo", alfa)
    tk = tukey_hsd(anova) if anova else pd.DataFrame()
    ls = lsd(anova) if anova else pd.DataFrame()
    return {
        "anova": anova, "tukey": tk, "lsd": ls,
        "letras_tukey": letras_cld(anova["medias"], tk) if anova else pd.DataFrame(),
        "letras_lsd": letras_cld(anova["medias"], ls) if anova else pd.DataFrame(),
    }

def _ordenar_columnas(tabla, columnas):
    return tabla[[c for c in columnas if c in tabla.columns]]

SI_DIFIEREN = "✔ Sí difieren"

NO_DIFIEREN = "✘ No difieren"

_FIGURAS = {}

def _es_columna_p(nombre):
    n = str(nombre).lower()
    return n.startswith("p-valor") or n.startswith("p ajustada") or n.startswith("p sin")

def _fmt(df, dec=3):
    """Formato de tabla: 3 decimales (4 en p-valores), enteros sin decimales y
    celdas vacías en blanco en lugar de 'None'."""
    formatos = {}
    for c in df.columns:
        if pd.api.types.is_float_dtype(df[c]):
            formatos[c] = "{:.4f}" if _es_columna_p(c) else "{:.%df}" % dec
    try:
        return df.style.format(formatos, na_rep="")
    except Exception:
        return df.round(dec)

def _es_control(nombre):
    t = str(nombre).lower()
    return "control" in t or "patrón" in t or "patron" in t

def _unidades(variable):
    if variable == "Germinó":
        return "%", "%²", "Germinación"
    return LONG, f"{LONG}²", variable

def columna_germinacion(df):
    return "Germinó (conteo)" if "Germinó (conteo)" in df.columns else "Germinó"

def germinacion_por_grupo(df, columna="Grupo"):
    c = columna_germinacion(df)
    d = df.dropna(subset=[c])
    return d.groupby(columna)[c].agg(["sum", "count"])

def tabla_ranking(desc, letras_tk, letras_lsd, u, germ=None):
    t = desc[["Grupo", "n", "Media", "Desviación estándar"]].copy()
    t["Promedio ± DE"] = [f"{m:.2f} ± {s:.2f}" if pd.notna(s) else f"{m:.2f}"
                          for m, s in zip(t["Media"], t["Desviación estándar"])]
    t["Letras Tukey"] = t["Grupo"].map(dict(zip(letras_tk["Grupo"], letras_tk["Letras"]))) if not letras_tk.empty else ""
    t["Letras LSD"] = t["Grupo"].map(dict(zip(letras_lsd["Grupo"], letras_lsd["Letras"]))) if not letras_lsd.empty else ""
    t = t.sort_values("Media", ascending=False).reset_index(drop=True)
    t.insert(0, "Lugar", range(1, len(t) + 1))
    columnas = ["Lugar", "Grupo", "n", "Promedio ± DE", "Letras Tukey", "Letras LSD"]
    if germ is not None and not germ.empty:
        t["% germinación"] = [
            f"{germ.loc[g, 'sum'] / germ.loc[g, 'count'] * 100:.1f} % "
            f"({int(germ.loc[g, 'sum'])}/{int(germ.loc[g, 'count'])})"
            if g in germ.index and germ.loc[g, "count"] else "—" for g in t["Grupo"]]
        columnas.insert(4, "% germinación")
    return t[columnas].rename(columns={"Promedio ± DE": f"Promedio ± DE ({u})"})

def tabla_tukey_clara(tk, u):
    if tk.empty:
        return tk
    t = tk.copy()
    t["Conclusión"] = np.where(t["Resultado"] == SIG, SI_DIFIEREN, NO_DIFIEREN)
    t = t.rename(columns={
        "Grupo 1": "Grupo A", "Grupo 2": "Grupo B", "n1": "n A", "n2": "n B",
        "Media 1": f"Promedio A ({u})", "Media 2": f"Promedio B ({u})",
        "|Diferencia|": f"|A − B| ({u})", "DHS": f"DHS ({u})",
        "Comparación |dif| vs DHS": "|A − B| vs DHS",
        "Comparación q": "q calculada vs q crítica",
        "k (medias)": "k (grupos)", "GL error": "GL del error",
    })
    cols = ["Grupo A", "Grupo B", "Conclusión", f"|A − B| ({u})", f"DHS ({u})", "|A − B| vs DHS",
            "q calculada", "q crítica", "q calculada vs q crítica", "k (grupos)", "GL del error",
            "p ajustada", f"Promedio A ({u})", f"Promedio B ({u})", "n A", "n B"]
    return t[[c for c in cols if c in t.columns]]

def tabla_lsd_clara(ls, u):
    if ls.empty:
        return ls
    t = ls.copy()
    t["Conclusión"] = np.where(t["Resultado"] == SIG, SI_DIFIEREN, NO_DIFIEREN)
    t = t.rename(columns={
        "Grupo 1": "Grupo A", "Grupo 2": "Grupo B", "n1": "n A", "n2": "n B",
        "Media 1": f"Promedio A ({u})", "Media 2": f"Promedio B ({u})",
        "|Diferencia|": f"|A − B| ({u})", "LSD": f"LSD ({u})",
        "Comparación |dif| vs LSD": "|A − B| vs LSD",
        "Comparación t": "t calculada vs t crítica", "GL error": "GL del error",
    })
    cols = ["Grupo A", "Grupo B", "Conclusión", f"|A − B| ({u})", f"LSD ({u})", "|A − B| vs LSD",
            "t calculada", "t crítica", "t calculada vs t crítica", "GL del error", "p-valor",
            f"Promedio A ({u})", f"Promedio B ({u})", "n A", "n B"]
    return t[[c for c in cols if c in t.columns]]

def tabla_anova_clara(anova, u2):
    t = anova["tabla"].copy()
    col_fc = [c for c in t.columns if c.startswith("F crítica")][0]
    t["Conclusión"] = t["Resultado"].map({SIG: "✔ Hay diferencia", NO_SIG: "✘ No hay diferencia"}).fillna("")
    t = t.rename(columns={"Fuente": "Fuente de variación",
                          "Suma de cuadrados": f"SC ({u2})", "Cuadrado medio": f"CM ({u2})",
                          col_fc: "F crítica", "Comparación": "F calc. vs F crít."})
    t = t.drop(columns=["Resultado"])
    return t[["Fuente de variación", "GL", f"SC ({u2})", f"CM ({u2})", "F calculada",
              "F crítica", "F calc. vs F crít.", "p-valor", "Conclusión"]]

def pares_significativos(tk, u):
    """Frases tipo 'A (4.72 cm) > B (3.43 cm)' para los pares que difieren."""
    frases = []
    for _, f in tk[tk["Resultado"] == SIG].iterrows():
        a, b, ma, mb = f["Grupo 1"], f["Grupo 2"], f["Media 1"], f["Media 2"]
        if mb > ma:
            a, b, ma, mb = b, a, mb, ma
        frases.append(f"**{a}** ({ma:.2f} {u}) es mayor que **{b}** ({mb:.2f} {u})")
    return frases

def texto_reporte(variable, u, agrupacion, unidad_txt, anova, tk, alfa, desc):
    nombre = "el porcentaje de germinación" if variable == "Germinó" else f"la longitud de {variable.lower()}"
    signo = "indicó" if anova["significativo"] else "no indicó"
    txt = (f"Se compararon los promedios de {nombre} ({u}) entre grupos de "
           f"{agrupacion.lower()}, usando {anova['n total']} {unidad_txt} como réplicas. "
           f"El ANOVA de una vía {signo} diferencias significativas "
           f"(F calculada = {anova['F calculada']:.3f}, F crítica = {anova['F crítica']:.3f}; "
           f"GL = {anova['GL entre']} y {anova['GL error']}; p = {anova['p-valor']:.4f}; α = {alfa}). ")
    sig = tk[tk["Resultado"] == SIG] if not tk.empty else tk
    if anova["significativo"] and not sig.empty:
        partes = []
        for _, f in sig.iterrows():
            a, b, ma, mb = f["Grupo 1"], f["Grupo 2"], f["Media 1"], f["Media 2"]
            if mb > ma:
                a, b, ma, mb = b, a, mb, ma
            partes.append(f"{a} ({ma:.2f} {u}) fue mayor que {b} ({mb:.2f} {u})")
        txt += f"Según la prueba de Tukey (α = {alfa}), " + "; ".join(partes) + ". "
    elif anova["significativo"]:
        txt += (f"Aunque el ANOVA fue significativo, la prueba de Tukey (α = {alfa}) no "
                "separó ningún par de grupos. ")
    mejor = desc.sort_values("Media", ascending=False).iloc[0]
    txt += f"El promedio más alto fue el de {mejor['Grupo']} ({mejor['Media']:.2f} {u})."
    return txt

COLOR_LINEA_PATRON = "#e8a33d"

def _valor_patron(grupos, medias, referencia):
    for g, m in zip(grupos, medias):
        if _es_control(g) and pd.notna(m):
            return float(m)
    if referencia is not None and referencia.get("n", 0) > 0 and pd.notna(referencia.get("Media")):
        return float(referencia["Media"])
    return None

def tabla_letras_grupos(letras):
    """Para cada letra, qué grupos la llevan (esos grupos no difieren entre sí)."""
    if letras is None or letras.empty:
        return pd.DataFrame()
    ordenado = letras.sort_values("Media", ascending=False)
    todas = sorted({c for l in ordenado["Letras"] for c in str(l)})
    filas = []
    for c in todas:
        grupos = [str(g) for g, l in zip(ordenado["Grupo"], ordenado["Letras"]) if c in str(l)]
        filas.append({"Letra": c, "Grupos que la llevan (no difieren entre sí)": " · ".join(grupos),
                      "Cuántos": len(grupos)})
    return pd.DataFrame(filas)

def _forma_bowley(b, n):
    if n < 4 or pd.isna(b):
        return "Pocos datos (n < 4)"
    if abs(b) < 0.1:
        return "Simétrica"
    grado = "ligeramente" if abs(b) < 0.3 else "muy"
    lado = "hacia valores altos" if b > 0 else "hacia valores bajos"
    return f"{grado.capitalize()} sesgada {lado}"

def tabla_cuartiles(datos, variable, factor=1.5):
    filas = []
    for g, sub in datos.groupby("Grupo"):
        x = pd.to_numeric(sub[variable], errors="coerce").dropna().astype(float)
        n = len(x)
        if n == 0:
            continue
        q1, med, q3 = x.quantile([0.25, 0.5, 0.75]).tolist()
        iqr = q3 - q1
        lim_inf, lim_sup = q1 - factor * iqr, q3 + factor * iqr
        dentro = x[(x >= lim_inf) & (x <= lim_sup)]
        bowley = (q3 + q1 - 2 * med) / iqr if iqr > 0 else np.nan
        filas.append({
            "Grupo": str(g), "n": n, "Mínimo": x.min(), "Q1 (25 %)": q1, "Mediana (Q2)": med,
            "Q3 (75 %)": q3, "Máximo": x.max(), "IQR (Q3 − Q1)": iqr,
            "Límite inferior": lim_inf, "Límite superior": lim_sup,
            "Atípicos bajos": int((x < lim_inf).sum()), "Atípicos altos": int((x > lim_sup).sum()),
            "Promedio": x.mean(), "Asimetría de Bowley": bowley,
            "Forma de la distribución": _forma_bowley(bowley, n),
            "_bigote_inf": dentro.min() if len(dentro) else q1,
            "_bigote_sup": dentro.max() if len(dentro) else q3,
        })
    t = pd.DataFrame(filas)
    if t.empty:
        return t
    return t.sort_values("Mediana (Q2)", ascending=False).reset_index(drop=True)

def diagnostico_cuartiles(tc, factor=1.5):
    validos = tc[(tc["n"] >= 4) & (tc["IQR (Q3 − Q1)"] > 0)]
    salida = {"razon_iqr": np.nan, "iqr_max": None, "iqr_min": None,
              "sesgados": [], "atipicos": int(tc["Atípicos bajos"].sum() + tc["Atípicos altos"].sum()),
              "grupos_validos": len(validos)}
    if len(validos) >= 2:
        mx = validos.loc[validos["IQR (Q3 − Q1)"].idxmax()]
        mn = validos.loc[validos["IQR (Q3 − Q1)"].idxmin()]
        salida.update(razon_iqr=mx["IQR (Q3 − Q1)"] / mn["IQR (Q3 − Q1)"],
                      iqr_max=(mx["Grupo"], mx["IQR (Q3 − Q1)"]),
                      iqr_min=(mn["Grupo"], mn["IQR (Q3 − Q1)"]))
    salida["sesgados"] = validos.loc[validos["Asimetría de Bowley"].abs() >= 0.3, "Grupo"].tolist()
    return salida

def pasos_cuartil(valores, p):
    x = np.sort(np.asarray(valores, dtype=float))
    n = len(x)
    pos = 1 + (n - 1) * p
    baja = int(math.floor(pos))
    frac = pos - baja
    v_baja = x[baja - 1]
    v_alta = x[min(baja, n - 1)]
    return pos, baja, frac, v_baja, v_alta, v_baja + frac * (v_alta - v_baja)

def _rgba(hex_color, alfa):
    h = hex_color.lstrip("#")
    return f"rgba({int(h[0:2], 16)},{int(h[2:4], 16)},{int(h[4:6], 16)},{alfa})"

NOMBRES_FIGURAS = {
    "res_graf": "Promedio por grupo con letras",
    "desc_dispersion": "Datos de cada grupo alrededor de su promedio",
    "anova_cuartiles": "Cuartiles por grupo",
    "tk_cluster": "Tukey: diferencias entre pares",
    "tk_barras": "Tukey: promedio de cada grupo con sus letras",
    "lsd_intervalos": "LSD: intervalos de comparación",
    "lsd_pares": "LSD: qué pares difieren",
    "mas_raiz_tallo": "Raíz y tallo por grupo",
    "mas_germinacion": "Germinaron vs no germinaron",
    "mas_correlacion": "Correlación entre raíz y tallo",
    "mas_corridas": "Comparación entre cajas",
}

_TOKEN_MARKDOWN = re.compile(r"\*\*(.+?)\*\*")

def _archivo_seguro(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_").lower()[:60] or "grafica"

CSS_DESCARGAS = """
*{box-sizing:border-box}
body{font-family:Arial,Helvetica,sans-serif;margin:0;background:#fff;color:#222;font-size:14px}
.barra{position:sticky;top:0;z-index:5;background:#f3f5f9;border-bottom:1px solid #d8dde6;
  padding:10px 14px;display:flex;flex-wrap:wrap;gap:8px;align-items:center}
.barra button,.una button{border:1px solid #2a78d6;background:#2a78d6;color:#fff;border-radius:6px;
  padding:7px 12px;font-size:13px;cursor:pointer}
.barra button.sec,.una button{background:#fff;color:#2a78d6}
.barra button:disabled{opacity:.5;cursor:wait}
.estado{font-size:13px;color:#555;margin-left:6px}
.ayuda{background:#fff8e6;border:1px solid #f0d58c;border-radius:6px;padding:8px 12px;margin:12px 18px;font-size:13px}
.vista{padding:6px 22px 30px}
h1{font-size:22px;color:#1f4e8c;margin:14px 0 2px}
.sub{color:#666;font-size:12.5px;margin:0 0 12px}
h2{font-size:16px;color:#1f4e8c;margin:20px 0 6px;border-bottom:1px solid #dfe5ee;padding-bottom:3px}
h3{font-size:14px;margin:10px 0 6px}
p{margin:4px 0 6px;line-height:1.45}
.nota{color:#666;font-size:12px}
ul{margin:4px 0 8px 18px;padding:0}
table{border-collapse:collapse;width:100%;font-size:12px;margin:6px 0 10px}
th{background:#2a78d6;color:#fff;padding:4px 6px;text-align:left;font-weight:600}
td{border:1px solid #d7dce5;padding:3px 6px}
tr:nth-child(even) td{background:#f3f6fb}
table.kv td:first-child{font-weight:600;width:30%}
figure{margin:10px 0 14px;page-break-inside:avoid;break-inside:avoid}
figure img{width:100%;border:1px solid #e3e7ee}
figcaption{font-weight:600;font-size:13px;margin-bottom:4px}
.pagina{margin:18px 0 26px}
.pagina img{width:100%;border:1px solid #e3e7ee}
.una{display:flex;gap:6px;margin-top:6px}
.oculto{display:none}
@media print{
  .barra,.ayuda,.una{display:none !important}
  body.modo-reporte #graficas{display:none}
  body.modo-graficas #reporte{display:none}
  .pagina{page-break-after:always;break-after:page;margin:0}
  .pagina:last-child{page-break-after:auto;break-after:auto}
  figure img,.pagina img{border:none}
}
@page{size:letter;margin:12mm}
"""

JS_DESCARGAS = """
const IMGS = {};
const estado = document.getElementById('estado');
const botones = () => document.querySelectorAll('.barra button, .una button');
function bloquear(si){ botones().forEach(b => b.disabled = si); }
async function imagen(i, fmt){
  const f = FIGS[i];
  if (fmt === 'png' && IMGS[i]) return IMGS[i];
  return Plotly.toImage(f.fig, {format: fmt, width: f.w, height: f.h, scale: 2});
}
function bajar(url, nombre){
  const a = document.createElement('a'); a.href = url; a.download = nombre;
  document.body.appendChild(a); a.click(); a.remove();
}
const espera = ms => new Promise(r => setTimeout(r, ms));
async function una(i, fmt){
  const url = await imagen(i, fmt);
  bajar(url, String(i + 1).padStart(2, '0') + '_' + FIGS[i].archivo + (fmt === 'jpeg' ? '.jpg' : '.png'));
}
async function todas(fmt){
  bloquear(true);
  for (let i = 0; i < FIGS.length; i++){
    estado.textContent = 'Descargando ' + (i + 1) + ' de ' + FIGS.length + '…';
    await una(i, fmt); await espera(450);
  }
  estado.textContent = 'Listo: ' + FIGS.length + ' imágenes descargadas.';
  bloquear(false);
}
function imprimir(modo){
  document.body.className = 'modo-' + modo;
  estado.textContent = 'En la ventana que se abre, elige “Guardar como PDF” en Destino.';
  setTimeout(() => window.print(), 150);
}
async function preparar(){
  if (typeof Plotly === 'undefined'){ estado.textContent = 'No se pudo cargar el dibujador de gráficas.'; return; }
  bloquear(true);
  for (let i = 0; i < FIGS.length; i++){
    estado.textContent = 'Preparando gráfica ' + (i + 1) + ' de ' + FIGS.length + '…';
    try {
      IMGS[i] = await imagen(i, 'png');
      document.querySelectorAll('[data-img="' + i + '"]').forEach(im => im.src = IMGS[i]);
    } catch (e) { console.log(e); }
  }
  estado.textContent = 'Listo. Elige qué descargar.';
  bloquear(false);
}
preparar();
"""

def preparar_patron_vs_corridas(df_semilla, dims, variable, por_caja, juntar_patrones=False):
    """Arma los grupos para comparar los patrones (Clase S, sin tratamiento) contra las corridas.
    Patrones: cada uno por separado (T-01, T-02…) o todos juntos. Corridas: agrupadas con
    'Comparar por'. Devuelve (datos de análisis, datos por semilla, por_caja_usado, lista patrones)."""
    d = df_semilla.copy()
    es_patron = d["Clase"].astype(str) == "S"
    if not es_patron.any() or es_patron.all():
        return None
    pat = d[es_patron]
    varias_var = pat["Variedad"].nunique() > 1
    if juntar_patrones:
        etiqueta = pd.Series("Patrón", index=pat.index)
    else:
        etiqueta = "Patrón " + pat["Bloque"].astype(str)
        if pat["Hoja"].nunique() > 1:
            etiqueta = etiqueta + " · " + pat["Hoja"].astype(str)
    if varias_var:
        etiqueta = etiqueta + " · " + pat["Variedad"].astype(str)
    d.loc[es_patron, "Grupo"] = etiqueta
    dims_corr = [x for x in dims if x != "Corrida"] or ["Tratamiento"]
    d.loc[~es_patron, "Grupo"] = construir_grupo(d[~es_patron], dims_corr).astype(str)
    d["Rol"] = np.where(es_patron, "Patrón", "Corrida")

    cajas = d.groupby("Grupo")["Corrida"].nunique()
    usar_caja = bool(por_caja and cajas.min() >= 2)
    if usar_caja:
        df_an = agregar_por_unidad(d, variable, COLUMNAS_CAJA + ["Rol"])
    else:
        df_an = d.dropna(subset=[variable]).copy()
        if variable == "Germinó":
            df_an[variable] = pd.to_numeric(df_an[variable], errors="coerce") * 100
    patrones = sorted(d.loc[es_patron, "Grupo"].unique())
    return df_an, d, usar_caja, patrones

def _veredicto(fila, corrida):
    """Texto corto para una comparación corrida vs patrón."""
    if fila is None:
        return "—"
    dif = fila["Media 1"] - fila["Media 2"] if fila["Grupo 1"] == corrida else fila["Media 2"] - fila["Media 1"]
    if fila["Resultado"] == SIG:
        return "▲ Mayor (✔ difiere)" if dif > 0 else "▼ Menor (✔ difiere)"
    return "＝ No difiere (✘)"

def matriz_corrida_vs_patron(desc, tk, ls, patrones, u):
    corridas = desc[~desc["Grupo"].isin(patrones)].sort_values("Media", ascending=False)
    media = dict(zip(desc["Grupo"], desc["Media"]))

    def buscar(tabla, a, b):
        if tabla is None or tabla.empty:
            return None
        m = ((tabla["Grupo 1"] == a) & (tabla["Grupo 2"] == b)) | ((tabla["Grupo 1"] == b) & (tabla["Grupo 2"] == a))
        return tabla[m].iloc[0] if m.any() else None

    filas = []
    for _, c in corridas.iterrows():
        fila = {"Corrida": c["Grupo"], f"Promedio ({u})": c["Media"]}
        for p in patrones:
            corto = p.replace("Patrón ", "")
            fila[f"{corto}: diferencia ({u})"] = c["Media"] - media.get(p, np.nan)
            fila[f"{corto}: Tukey"] = _veredicto(buscar(tk, c["Grupo"], p), c["Grupo"])
            fila[f"{corto}: LSD"] = _veredicto(buscar(ls, c["Grupo"], p), c["Grupo"])
        filas.append(fila)
    return pd.DataFrame(filas)

def _semillas_para_excel(crudo, analizado, tabla_cajas, variable):
    """Una fila por semilla, con las columnas que usa el programa para los cálculos."""
    info = tabla_cajas.set_index("_id")
    d = crudo.copy()
    d["Apartado"] = d["Corrida"].map(info["Apartado"])
    d["Caja"] = d["Corrida"].map(info["Caja"])
    d["_orden"] = d["Corrida"].map({c: i for i, c in enumerate(tabla_cajas["_id"])})
    d = d.sort_values(["_orden", "No"])
    usados = analizado.set_index("ID semilla")[variable] if variable in analizado else pd.Series(dtype=float)
    if variable == "Germinó":
        d["Entra al análisis"] = np.where(d["ID semilla"].isin(usados.index), "Sí", "No (descartada)")
    else:
        valor = d["ID semilla"].map(usados)
        d[f"{variable} usado en el análisis ({LONG})"] = valor
        d["Entra al análisis"] = np.where(valor.notna(), "Sí", "No (descartada)")
    columnas = {"Archivo": "Archivo", "Hoja": "Hoja / fecha", "Apartado": "Apartado", "Caja": "Caja",
                "Bloque": "Bloque", "Réplica": "Réplica", "Grupo": "Grupo comparado",
                "Variedad": "Variedad", "Clase nombre": "Parte expuesta", "Temperatura": "Temperatura",
                "Tratamiento": "Tiempo de tratamiento", "No": "Semilla n.º", "Germinó": "Germinó (1 = sí)",
                "Raíz": f"Raíz ({LONG})", "Tallo": f"Tallo ({LONG})"}
    extra = [c for c in d.columns if c.endswith(f"usado en el análisis ({LONG})")] + ["Entra al análisis"]
    vista = d[[c for c in columnas if c in d.columns] + extra].rename(columns=columnas)
    return vista, d.drop(columns=["_orden"])

def _excel_crudo(vista, completo, tabla_cajas, variable):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        vista.to_excel(xw, sheet_name="Semillas", index=False)
        tabla_cajas.drop(columns=["_id", "_sel"]).to_excel(xw, sheet_name="Resumen por caja", index=False)
        completo.to_excel(xw, sheet_name="Todas las columnas", index=False)
        pd.DataFrame({"Nota": [
            f"Semillas: una fila por semilla tal como la lee el programa (longitudes en {LONG}).",
            f"'Entra al análisis' y el valor usado se refieren a la variable elegida ({variable}) "
            "con los descartes aplicados en el dashboard.",
            "Resumen por caja: la misma tabla de la pestaña Caja por caja.",
            "Todas las columnas: el registro interno completo con el que se hacen los cálculos.",
        ]}).to_excel(xw, sheet_name="Léeme", index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, ancho + 2), 45)
            ws.freeze_panes = "A2"
    return buf.getvalue()

def _excel_como_original(completo, tabla_cajas, archivo):
    """Un bloque por caja como en el Excel que se sube: título, encabezados No / Raíz / Tallo /
    Germinó y una semilla por fila. Una hoja por fecha/hoja; patrón primero y luego corridas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    wb.remove(wb.active)
    negrita, centro = Font(bold=True), Alignment(horizontal="center")
    fondo_ap = {"Patrón": PatternFill("solid", fgColor="E7E6E1"), "Corrida": PatternFill("solid", fgColor="DCE9F7")}
    linea = Side(style="thin", color="999999")
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    usados = set()
    orden = {c: i for i, c in enumerate(tabla_cajas["_id"])}
    for hoja, sub in completo.groupby("Hoja", sort=False):
        nombre = re.sub(r"[\\/*?:\[\]]", "-", str(hoja))[:31] or "Hoja"
        base, k = nombre, 2
        while nombre in usados:
            nombre = f"{base[:28]} ({k})"; k += 1
        usados.add(nombre)
        ws = wb.create_sheet(nombre)
        ws["A1"] = f"Longitud dada en {LONG} · exportado del dashboard BioGrow · {archivo} · {hoja}"
        ws["A1"].font = Font(italic=True, color="555555")
        col = 1
        for caja in sorted(sub["Corrida"].unique(), key=lambda c: orden.get(c, 1e9)):
            c = sub[sub["Corrida"] == caja].sort_values("No")
            ap = str(c["Apartado"].iloc[0])
            ap_corto = "Patrón" if ap.startswith("Patrón") else "Corrida"
            ws.cell(row=2, column=col, value=ap).font = Font(bold=True, color="444444")
            titulo = ws.cell(row=3, column=col, value=str(c["Bloque"].iloc[0]))
            titulo.font, titulo.fill = Font(bold=True, size=12), fondo_ap[ap_corto]
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 3)
            titulo.alignment = centro
            for j, enc in enumerate(["No", "Raíz", "Tallo", "Germinó"]):
                h = ws.cell(row=4, column=col + j, value=enc)
                h.font, h.alignment, h.border, h.fill = negrita, centro, borde, fondo_ap[ap_corto]
            for i, (_, f) in enumerate(c.iterrows()):
                for j, campo in enumerate(["No", "Raíz", "Tallo", "Germinó"]):
                    v = f[campo]
                    v = None if pd.isna(v) else (int(v) if campo in ("No", "Germinó") else float(v))
                    celda = ws.cell(row=5 + i, column=col + j, value=v)
                    celda.border = borde
                    if campo in ("Raíz", "Tallo"):
                        celda.number_format = "0.00"
            for j in range(4):
                ws.column_dimensions[ws.cell(row=4, column=col + j).column_letter].width = 9
            ws.column_dimensions[ws.cell(row=4, column=col + 4).column_letter].width = 3
            col += 5
        ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
